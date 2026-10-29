"""Notas de precios: tasas en porcentaje e incidencias en puntos porcentuales."""
from docx.shared import Inches
from notas_layout import Layout
from notas_imoam import period_text

def selected(f):
    up=sorted([r for r in f['generics'] if r['incidence']>0],key=lambda r:-r['incidence'])[:3]
    down=sorted([r for r in f['generics'] if r['incidence']<0],key=lambda r:r['incidence'])[:3]
    return up+down

def hero(f):
    r=f['rows'][0]
    return [[f['indicator']+' · '+period_text(f['ym']).capitalize(),'Publicación '+f['publication_date']],
            [f'Índice {f["index"]:.3f}' if f['index'] is not None else 'Total incluido petróleo',f'Mensual {r["mom"]:.2f}% | Anual {r["yoy"]:.2f}%']]

def core(f):
    ids=['total','subyacente','mercancias','servicios','no_subyacente','agro','energia_tarifas'] if f['indicator']=='INPC' else [r['id'] for r in f['rows']]
    return [next(r for r in f['rows'] if r['id']==i) for i in ids]

def tables(f):
    return [hero(f),[['Componente','Mensual (%)','Anual (%)']]+[[r['label'],f'{r["mom"]:.2f}',f'{r["yoy"]:.2f}'] for r in core(f)],
            [['Producto genérico','Mensual (%)','Incidencia (pp)']]+[[r['label'],f'{r["mom"]:.2f}',f'{r["incidence"]:.3f}'] for r in selected(f)]]

def bars(x,name,labels,values,title,unit,digits=2):
    import matplotlib.pyplot as plt
    with plt.rc_context({'font.family':'Noto Sans','font.size':9,'axes.unicode_minus':False}):
        fig,ax=plt.subplots(figsize=(6.1,1.85 if len(values)>4 else 1.65));ys=list(range(len(values)))
        ax.barh(ys,values,color=['#611232' if v>=0 else '#235B4C' for v in values]);ax.axvline(0,color='#888888',lw=.7)
        ax.set_yticks(ys,labels);ax.invert_yaxis();ax.set_title(title,fontsize=11,color='#10312B');ax.set_xlabel(unit)
        pad=max(.001,max(abs(v) for v in values))*.3;ax.set_xlim(min(0,min(values))-pad,max(0,max(values))+pad)
        for y,v in enumerate(values):
            t=f'{v:.{digits}f}';ax.text(v,y,' '+t if v>=0 else t+' ',ha='left' if v>=0 else 'right',va='center',fontsize=8)
        ax.spines[['top','right']].set_visible(False);fig.tight_layout();path=x.out.parent/(name+'.png');fig.savefig(path,dpi=300);plt.close(fig)
    x.doc.add_picture(str(path),width=Inches(6.05))

def build(f,url,template,out):
    x=Layout(template,out);key=f['indicator'];r=f['rows'][0];data=tables(f)
    title='Precios al consumidor' if key=='INPC' else 'Precios al productor'
    x.table(data[0],True);x.p(title,bold=True,color='10312B');x.note('Base: '+f['base']+'. Variaciones de precios publicadas en el boletín; no son variaciones reales de producción.')
    x.h('Resumen ejecutivo')
    diff=r['yoy']-r['prior_yoy']
    x.p(f'En {period_text(f["ym"])}, la variación anual del {key} fue {r["yoy"]:.2f}%, frente a {r["prior_yoy"]:.2f}% en el mismo mes del año anterior: una diferencia de {diff:.2f} puntos porcentuales.',gray=True)
    if key=='INPC':
        sub=next(v for v in f['rows'] if v['id']=='subyacente');nonsub=next(v for v in f['rows'] if v['id']=='no_subyacente')
        x.bullet(f'La inflación subyacente registró {sub["yoy"]:.2f}% anual y la no subyacente, {nonsub["yoy"]:.2f}%. Sus incidencias mensuales fueron {sub["incidence"]:.3f} y {nonsub["incidence"]:.3f} puntos porcentuales.',True)
        x.bullet(f'El nivel general de precios aumentó {r["mom"]:.2f}% en el mes. Una reducción de la inflación anual no implica una caída del nivel de precios.',True)
    else:
        inter=f['groups'][1][0];final=f['groups'][2][0]
        x.bullet(f'Los precios de uso intermedio variaron {inter["mom"]:.2f}% mensual y {inter["yoy"]:.2f}% anual; los finales, {final["mom"]:.2f}% y {final["yoy"]:.2f}%, respectivamente. Incluyen petróleo.',True)
        energy=next(v for v in f['rows'] if v['id']=='energia')
        x.bullet(f'El componente de energía eléctrica, agua y gas registró {energy["yoy"]:.2f}% anual. Este resultado sectorial no representa el comportamiento de todos los costos productivos.',True)
    x.h('Variaciones por componente');x.table(data[1]);x.note('Fuente: INEGI, cuadro 1, p. 2. Comparaciones mensuales y anuales del mismo mes; no se mezclan quincenas ni se suman tasas de componentes.')
    ids=['mercancias','servicios','agro','energia_tarifas'] if key=='INPC' else ['primarias','secundarias','terciarias']
    rr=[next(v for v in f['rows'] if v['id']==i) for i in ids]
    bars(x,'anual',[v['label'] for v in rr],[v['yoy'] for v in rr],'Variación anual de los precios por componente','Variación anual (%)')
    x.h('Lectura del indicador')
    x.p('El INPC resume precios de una canasta representativa del consumo de los hogares. Las diferencias entre rubros muestran comportamientos distintos; no equivalen al cambio de precios experimentado por cada hogar.' if key=='INPC' else 'El INPP cubre precios de la producción nacional para consumo interno y exportación. Su variación no mide volumen de producción ni anticipa un traslado automático o inmediato a los precios al consumidor.')
    x.doc.add_page_break();x.h('Productos con incidencia al alza y a la baja');x.table(data[2]);x.note('Selección: tres mayores incidencias positivas y tres negativas entre los genéricos destacados por el INEGI. No es un listado exhaustivo de toda la canasta. Incidencia mensual en puntos porcentuales (pp).')
    rr=selected(f);labels=[v['label'].replace('Paquetes de internet, telefonía y televisión de paga','Internet, telefonía y TV') for v in rr]
    bars(x,'incidencias',labels,[v['incidence'] for v in rr],'Contribuciones publicadas a la variación mensual','Incidencia mensual (puntos porcentuales)',3)
    x.note('Fuente: '+('cuadro 2, p. 4' if key=='INPC' else 'cuadro 3, p. 5')+'. Las barras usan incidencias oficiales; no se derivan multiplicando tasas por ponderaciones supuestas.')
    x.h('Interpretación de las incidencias')
    x.p('La variación porcentual describe el cambio del precio de cada producto. La incidencia mide su contribución al movimiento del índice general: una variación mayor no implica necesariamente una mayor incidencia. Esta selección no reconstruye la variación total.')
    if key=='INPP':
        a=next(v for v in f['destinations'] if v['id']=='interna');b=next(v for v in f['destinations'] if v['id']=='exportacion')
        x.p(f'Por destino, los precios de bienes y servicios finales de demanda interna registraron {a["yoy"]:.2f}% anual, y los de exportación, {b["yoy"]:.2f}%. Son variaciones de precios, no de valores o volúmenes exportados.')
    x.h('Notas metodológicas')
    if key=='INPC':
        x.note('La comparación anual anterior corresponde al mismo mes de un año antes; no al mes inmediato anterior. Las incidencias pueden presentar diferencias de redondeo respecto al total. El boletín contiene niveles quincenales, finalidades del consumo y aperturas territoriales; esta nota utiliza la lectura mensual nacional.')
    else:
        x.note('Todas las cifras de esta nota corresponden a la cobertura incluida petróleo. No se equiparan con la serie sin petróleo. Los resultados del INPP son preliminares y susceptibles de actualización hasta cinco meses después de su publicación, conforme a la metodología de esta edición. No se atribuye causalidad sin evidencia.')
    x.p('',6.5);x.note(f'INEGI, boletín {f["bulletin"]}, publicado el {f["publication_date"]}. Lectura completa de {len(f["pages"])} páginas; metodología en la última página. Variaciones e incidencias de la misma edición.');x.note(url);x.save(title+' · '+f['period'])

def validate(f,doc):
    expected=tables(f)
    if len(doc.tables)!=3 or len(doc.inline_shapes)!=2:raise ValueError('Faltan tablas o gráficas')
    for actual,wanted in zip(doc.tables,expected):
        if [[c.text for c in r.cells] for r in actual.rows]!=wanted:raise ValueError('El Word difiere del boletín')
