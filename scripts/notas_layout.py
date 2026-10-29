"""Elementos Word nativos para notas institucionales de dos páginas."""
from docx import Document
from docx.shared import Pt,Inches,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH as ALIGN
import notas_word as W

class Layout:
    def __init__(self,template,out):
        W.fonts();self.out=out;self.doc=Document(template)
        for n in list(self.doc.element.body):
            if n.tag!=qn('w:sectPr'):self.doc.element.body.remove(n)
        for st in self.doc.styles:
            if st.type==1:st.font.name='Noto Sans';st.font.size=Pt(11)
        ns=self.doc.part.numbering_part.element
        aid=max([int(n.get(qn('w:abstractNumId'))) for n in ns.findall(qn('w:abstractNum'))]+[-1])+1
        self.nid=max([int(n.get(qn('w:numId'))) for n in ns.findall(qn('w:num'))]+[0])+1
        a=OxmlElement('w:abstractNum');a.set(qn('w:abstractNumId'),str(aid));lvl=OxmlElement('w:lvl');lvl.set(qn('w:ilvl'),'0')
        for tag,val in [('numFmt','bullet'),('lvlText','•'),('suff','space')]:
            x=OxmlElement('w:'+tag);x.set(qn('w:val'),val);lvl.append(x)
        a.append(lvl);ns.append(a);n=OxmlElement('w:num');n.set(qn('w:numId'),str(self.nid));x=OxmlElement('w:abstractNumId');x.set(qn('w:val'),str(aid));n.append(x);ns.append(n)

    def p(self,text='',size=11,bold=False,color='000000',keep=False,gray=False):
        p=self.doc.add_paragraph();p.alignment=ALIGN.JUSTIFY
        pf=p.paragraph_format;pf.space_before=Pt(0);pf.space_after=Pt(0);pf.line_spacing=1;pf.keep_with_next=keep
        r=p.add_run(text);r.font.name='Noto Sans';r.font.size=Pt(size);r.bold=bold;r.font.color.rgb=RGBColor.from_string(color)
        if gray:
            sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'F1F2F0');p._p.get_or_add_pPr().append(sh)
        return p

    def h(self,text):self.p('',6.5);self.p(text,bold=True,color='235B4C',keep=True)
    def note(self,text):self.p(text,8,color='666666')
    def bullet(self,text,gray=False):
        p=self.p(text,gray=gray);pr=OxmlElement('w:numPr')
        for tag,val in [('ilvl','0'),('numId',str(self.nid))]:
            x=OxmlElement('w:'+tag);x.set(qn('w:val'),val);pr.append(x)
        p._p.get_or_add_pPr().append(pr)

    def table(self,data,hero=False):
        t=self.doc.add_table(rows=0,cols=len(data[0]));t.autofit=False
        widths={2:[3.05,3.05],3:[3.1,1.5,1.5],4:[2.65,1.15,1.15,1.15]}[len(data[0])]
        for col,w in zip(t.columns,widths):col.width=Inches(w)
        for i,row in enumerate(data):
            rr=t.add_row();rr._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
            if i==0:
                repeat=OxmlElement('w:tblHeader');rr._tr.get_or_add_trPr().append(repeat)
            for cell,text,w in zip(rr.cells,row,widths):
                cell.width=Inches(w);cell.text=str(text);sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'10312B' if hero or i==0 else ('F4F4F2' if i%2 else 'FFFFFF'));cell._tc.get_or_add_tcPr().append(sh)
                for p in cell.paragraphs:
                    p.paragraph_format.space_before=Pt(0);p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1
                    for r in p.runs:r.font.name='Noto Sans';r.font.size=Pt(10);r.bold=hero or i==0;r.font.color.rgb=RGBColor.from_string('FFFFFF' if hero or i==0 else '000000')
        return t

    def bars(self,name,labels,values,title,unit,other=None,legend=None):
        import matplotlib.pyplot as plt
        with plt.rc_context({'font.family':'Noto Sans','font.size':9,'axes.unicode_minus':False}):
            fig,ax=plt.subplots(figsize=(6.1,1.75));ys=list(range(len(labels)))
            if other is None:
                ax.barh(ys,values,color=['#235B4C' if v>=0 else '#611232' for v in values])
                pairs=[(ys,values)]
            else:
                a=[y-.17 for y in ys];b=[y+.17 for y in ys]
                ax.barh(a,values,height=.31,color='#235B4C');ax.barh(b,other,height=.31,color='#A57F2C');pairs=[(a,values),(b,other)]
            allvals=list(values)+(list(other) if other is not None else [])
            pad=max(1,max(abs(v) for v in allvals))*.25;ax.set_xlim(min(0,min(allvals))-pad,max(0,max(allvals))+pad)
            for ypos,vals in pairs:
                for y,v in zip(ypos,vals):ax.text(v,y,f' {v:.1f}' if v>=0 else f'{v:.1f} ',ha='left' if v>=0 else 'right',va='center',fontsize=8)
            ax.set_yticks(ys,labels);ax.invert_yaxis();ax.axvline(0,color='#777777',lw=.7);ax.set_title(title,fontsize=11,color='#10312B');ax.set_xlabel(unit,fontsize=9)
            ax.spines[['top','right']].set_visible(False);fig.tight_layout();path=self.out.parent/(name+'.png');fig.savefig(path,dpi=300);plt.close(fig)
        self.doc.add_picture(str(path),width=Inches(6.05))
        if legend:self.note(legend)

    def save(self,title):self.doc.core_properties.title=title;self.doc.save(self.out)
