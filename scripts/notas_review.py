"""Controles locales reproducibles. No usa modelos, claves ni servicios de IA.

La composición se comprueba por estructura y geometría; no equivale a una
revisión visual humana. Los cambios de metodología requieren intervención.
"""
from __future__ import annotations
import csv
import json
import re
from pathlib import Path

import pdfplumber
from docx import Document
from docx.oxml.ns import qn
from PIL import Image
import notas_imoam as I

VERSION='reglas-python-1'
KEYS=('datos','logica','estructura','render')


def validate(review: dict) -> bool:
    if set(review)!=set((*KEYS,'issues')) or not isinstance(review.get('issues'),list):
        raise ValueError('Revisión incompleta o esquema incorrecto')
    if not all(type(review[k]) is bool for k in KEYS) or not all(isinstance(x,str) for x in review['issues']):
        raise ValueError('Tipos incorrectos en la revisión')
    return all(review[k] for k in KEYS) and not review['issues']


def normalized(text):return re.sub(r'\s+',' ',text).strip()

def review(facts: dict,draft: str,rules: str,images: list[tuple[str,Path]], *_unused) -> dict:
    result={k:True for k in KEYS};result['issues']=[]
    def check(layer,condition,message):
        if not condition:result[layer]=False;result['issues'].append(message)
    notes=[p for label,p in images if label.startswith('Nota Word página')]
    sources=[p for label,p in images if label.startswith('Boletín oficial página')]
    if not notes:raise ValueError('Faltan las páginas de la nota')
    work=notes[0].parent
    # Segunda lectura del PDF y cotejo de la evidencia archivada.
    extracted=I.extract((work/'boletin.pdf').read_bytes())
    for key in ('values','ym','period','publication_date','source_sha256','context','base','bulletin'):
        check('datos',extracted[key]==facts[key],f'Hechos incompatibles con fuente: {key}')
    stored=json.loads((work/'hechos.json').read_text())
    check('datos',all(stored.get(k)==v for k,v in facts.items()),'El archivo de hechos difiere de la extracción')
    with (work/'trazabilidad.csv').open(encoding='utf-8-sig',newline='') as f: traced=list(csv.DictReader(f))
    check('datos',traced==I.trace(facts,stored['source_url']),'Trazabilidad incompleta o alterada')
    v=facts['values'];c=facts['context']
    check('logica',v[1]<=v[0]<=v[2] and v[4]<=v[3]<=v[5],'Estimación fuera de sus límites')
    check('logica',0<=c['coverage']<=100 and 0<c['days_ahead']<c['imai_lag'],'Cobertura o plazos incoherentes')
    check('logica',c['evaluation_start']<=c['evaluation_end']<facts['ym'],'Ventana histórica incompatible con el periodo')
    if v[4]<=0<=v[5]:check('logica','incluye cero' in draft and ('incertidumbre' in draft or 'no implica certeza' in draft),'Falta el contrapeso del intervalo que cruza cero')
    check('logica','desempeño pasado' in draft and 'no es la probabilidad' in draft,'Falta distinguir cobertura histórica de certeza actual')
    check('logica',bool(rules.strip()),'Faltan reglas editoriales versionadas')
    doc=Document(work/'IMOAM_nota.docx')
    native=' '.join(t.text or '' for t in doc.element.body.iter(qn('w:t')))
    check('estructura',normalized(native)==normalized(draft),'El contenido Word no coincide con el texto validado')
    check('estructura',len(doc.tables)==4 and len(doc.inline_shapes)==2,'Faltan tablas o visuales requeridos')
    check('estructura',len(doc.element.body.findall('.//'+qn('w:numPr')))>=5,'Faltan bullets nativos')
    headers=['Resumen ejecutivo','Estimación puntual e intervalos de confianza','Evaluación histórica y oportunidad de publicación','Cobertura y construcción del indicador','Notas de interpretación']
    check('estructura',all(x in native for x in headers),'Faltan secciones requeridas')
    expected=[[I.number(v[0]),I.number(v[1]),I.number(v[2])],[I.number(v[3],True),I.number(v[4],True),I.number(v[5],True)]]
    actual=[[cell.text for cell in row.cells[1:]] for row in doc.tables[1].rows[1:]] if len(doc.tables)>1 else []
    check('datos',actual==expected,'El cuadro Word difiere de las seis cifras oficiales')
    check('estructura',all(sec.header.part.blob and sec.footer.part.blob for sec in doc.sections),'Falta membrete')
    check('render',len(notes)==2 and len(sources)==len(facts['pages']),'Faltan páginas renderizadas de nota o boletín')
    for path in notes+sources:
        with Image.open(path) as im:
            im.load();check('render',im.width>=800 and im.height>=1000,'Render con resolución insuficiente')
    with pdfplumber.open(work/'IMOAM_nota.pdf') as pdf:
        check('render',len(pdf.pages)==2,'La nota debe tener exactamente dos páginas')
        rendered=' '.join(p.extract_text() or '' for p in pdf.pages)
        check('render',all(normalized(x) in normalized(rendered) for x in headers),'Sección ausente en el render')
        for page in pdf.pages:
            check('render',len(page.extract_text() or '')>=500,'Página sin contenido sustantivo')
            check('render',not any(ch['x0']<0 or ch['x1']>page.width+1 or ch['top']<0 or ch['bottom']>page.height+1 for ch in page.chars),'Texto fuera de página')
    return {'passed':validate(result),'result':result,'kind':'reglas_python','model':VERSION,
            'limits':['Controles de estructura y geometría; sin revisión visual humana ni análisis editorial por IA.']}
