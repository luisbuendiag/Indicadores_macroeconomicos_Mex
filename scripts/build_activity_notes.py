"""Notas de actividad desde boletines completos; sin servicios de IA."""
import argparse,csv,hashlib,json,os,re,urllib.request
from pathlib import Path
from urllib.parse import urlparse
from docx import Document
import notas_actividad as N
import notas_actividad_word as V
import notas_word as W
import build_auto_notes as A
import lib_notes_state as S

ROOT=Path(__file__).resolve().parents[1]
VERSION='actividad-python-1'

def allowed(url,key):
    u=urlparse(url)
    return u.scheme=='https' and u.hostname=='www.inegi.org.mx' and not u.username and not u.password and u.port in (None,443) and re.fullmatch(r'/contenidos/saladeprensa/boletines/\d{4}/'+key.lower()+r'/'+key.lower()+r'\d{4}_\d{2}\.pdf',u.path) is not None

def download(url,key):
    if not allowed(url,key):raise ValueError('Fuente no permitida')
    class Redirect(urllib.request.HTTPRedirectHandler):
        def redirect_request(self,req,fp,code,msg,headers,newurl):
            if not allowed(newurl,key):raise ValueError('Redirección no permitida')
            return super().redirect_request(req,fp,code,msg,headers,newurl)
    with urllib.request.build_opener(Redirect()).open(url,timeout=45) as r:data=r.read(20_000_001)
    if len(data)>20_000_000 or not data.startswith(b'%PDF'):raise ValueError('PDF inválido')
    return data

def validate_tables(f,doc):
    if len(doc.tables)!=3 or len(doc.inline_shapes)!=2:raise ValueError('Faltan tablas o gráficas')
    for table,fields in [(doc.tables[1],('mom','yoy')),(doc.tables[2],('original','ytd'))]:
        expected=[[f'{r[v]:.1f}' for v in fields] for r in f['rows']]
        actual=[[c.text for c in row.cells[1:]] for row in table.rows[1:]]
        if actual!=expected:raise ValueError('El cuadro Word difiere del boletín')

def run(key,root=ROOT,publish_requested=False,pdf_path=None):
    payload=json.loads((root/'data/indicadores.json').read_text());ind=payload['indicators'][key]
    states=S.load_state(root/'data/notas/estado.json');previous=states.get(key,{})
    state=dict(status='pendiente',target_period=ind['last_observation'],published=previous.get('published',{}),message='')
    work=None
    try:
        cfg=json.loads((root/'config/notas_automaticas.json').read_text())
        if pdf_path and publish_requested:raise ValueError('No publicar desde un PDF local')
        url=ind['url_boletin_oficial']
        if not allowed(url,key):raise ValueError('Boletín no oficial')
        pdf=pdf_path.read_bytes() if pdf_path else download(url,key)
        f=N.extract(pdf,key)
        if f['period']!=ind['last_observation']:raise ValueError('Periodo distinto al de la ficha')
        if int(f['publication_date'][:4])!=cfg['template_year']:raise ValueError('Actualizar el membrete institucional')
        snapshot=ind.get('boletin_cifras',{}).get(f['ym'],{})
        col=4 if key=='IMAI' else 1;expected=f['rows'][0]['original' if key=='IMAI' else 'mom']/100
        if str(col) in snapshot and snapshot[str(col)]['url']==url and abs(snapshot[str(col)]['value']-expected)>1e-8:raise ValueError('La fotografía del boletín difiere de su PDF')
        h=hashlib.sha256(pdf+(root/cfg['template']).read_bytes()+VERSION.encode())
        for name in ('build_activity_notes.py','notas_actividad.py','notas_actividad_word.py','notas_word.py','build_auto_notes.py'):h.update((root/'scripts'/name).read_bytes())
        fp=h.hexdigest();state['candidate_fingerprint']=fp
        dest=root/f'downloads/indicadores/{key}/nota/{key}_nota.docx';pub=previous.get('published',{})
        if pub.get('fingerprint')==fp and dest.exists() and S.digest(dest)==pub.get('docx_sha256'):return {**previous,'unchanged':True}
        work=Path(os.environ.get('NOTES_WORK_DIR',str(root/'.notes-work')))/key/fp
        work.mkdir(parents=True,exist_ok=True)
        (work/'boletin.pdf').write_bytes(pdf);(work/'boletin.txt').write_text('\n\n'.join(f['pages']))
        A.write_json(work/'hechos.json',{**f,'source_url':url})
        with (work/'trazabilidad.csv').open('w',encoding='utf-8-sig',newline='') as stream:
            writer=csv.DictWriter(stream,fieldnames=['variable','valor','unidad','periodo','serie','ubicacion','fuente']);writer.writeheader()
            for row in f['rows']:
                for field in ('index','mom','yoy','original','ytd'):
                    if row[field] is not None:writer.writerow(dict(variable=row['label']+' '+field,valor=row[field],unidad='puntos' if field=='index' else 'porcentaje',periodo=f['ym'],serie='original' if field in ('original','ytd') else 'desestacionalizada',ubicacion='página 4, cuadro 2' if field in ('original','ytd') else 'páginas 1–2, cuadro 1',fuente=url))
        out=work/f'{key}_nota.docx';V.build(f,url,root/cfg['template'],out)
        _,pages=W.render(out,2)
        if len(pages)!=2:raise ValueError('La nota no tiene dos páginas')
        validate_tables(N.extract((work/'boletin.pdf').read_bytes(),key),Document(out))
        result=dict(passed=True,kind='reglas_python',model=VERSION,result=dict(datos=True,logica=True,estructura=True,render=True,issues=[]),limits=['Controles locales; no equivalen a revisión visual humana.'])
        A.write_json(work/'revision.json',result)
        if publish_requested and os.environ.get('NOTES_AUTO_PUBLISH','true')=='true':return A.publish(root,work,f,fp,url,result,previous)
        state.update(status='validada_sin_publicar',message='Nota validada por reglas Python; publicación deshabilitada')
    except Exception as e:state.update(status='error',message='No se publicó la candidata: '+(str(e)[:180] if isinstance(e,(ValueError,RuntimeError)) else type(e).__name__))
    if work:A.write_json(work/'estado.json',state)
    if publish_requested:
        states[key]=state;A.write_json(root/'data/notas/estado.json',states);S.decorate(ind,root);A.write_json(root/'data/indicadores.json',payload)
    return state

if __name__=='__main__':
    p=argparse.ArgumentParser(description=__doc__);p.add_argument('--publish',action='store_true');p.add_argument('--key',choices=list(N.NAMES));p.add_argument('--pdf',type=Path);args=p.parse_args()
    if args.pdf and (args.publish or not args.key):p.error('--pdf requiere --key y no permite --publish')
    for key in [args.key] if args.key else N.NAMES:print(json.dumps({'indicator':key,**run(key,publish_requested=args.publish,pdf_path=args.pdf)},ensure_ascii=False))
