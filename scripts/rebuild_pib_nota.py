#!/usr/bin/env python3
"""Reconstrucción nativa de la nota PIB Oportuno (2T-2026).

Contexto: la edición original sólo existía como PDF (PIBOT_JUL_2026.pdf) y una
conversión mecánica pdf2docx produjo un archivo inválido para uso institucional
(compatibilityMode 14 + w:useFELayout, página rasterizada completa como imagen,
retícula rota, fuente Cambria). Este script reconstruye el documento como DOCX
nativo:

  * base estructural: _PLANTILLA_MAESTRA.docx (encabezado/pies/retícula/márgenes)
  * tipografía:      Noto Sans en todo el documento; cuerpo 11 pt
  * gráficas:        PNG a 200 dpi recortados del boletín oficial, guardados en
                     data/source/notas_machote/assets/pib/ (assets del machote)
  * salida:          PIB_machote.docx (estructura congelada) y
                     downloads/indicadores/PIB/nota/PIB_nota.docx (vigente)

No modifica cifras ni redacción; es estrictamente formato.
"""
from __future__ import annotations

import shutil
import sys
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Cm, Pt, RGBColor, Twips

ROOT = Path(__file__).resolve().parents[1]
PLANTILLA = ROOT / "data/source/notas_machote/_PLANTILLA_MAESTRA.docx"
MACHOTE = ROOT / "data/source/notas_machote/PIB_machote.docx"
NOTA = ROOT / "downloads/indicadores/PIB/nota/PIB_nota.docx"
ASSETS = ROOT / "data/source/notas_machote/assets/pib"

FONT = "Noto Sans"
GREEN_DARK = "10312B"   # barra delgada / texto principal sobre claro
GREEN = "235B4C"        # banda institucional
CELL_BG = "F4F4F2"      # celdas KPI / tabla
WINE = "670E35"
BORDER = "BFBFBF"
CONTENT_W = Twips(8838)  # ancho útil de la plantilla (carta, márgenes 1701)


def _set_run(run, size=11, bold=False, color=None, italic=False):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.insert(0, rFonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rFonts.set(qn(attr), FONT)


def _shade_cell(cell, hex_color):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hex_color)
    cell._tc.get_or_add_tcPr().append(shd)


def _cell_para(cell, lines, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=0):
    """lines: lista de (texto, size, bold, color)."""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(space_after)
    for i, (txt, size, bold, color) in enumerate(lines):
        r = p.add_run(txt)
        _set_run(r, size=size, bold=bold, color=color)
        if i < len(lines) - 1:
            r.add_break()
    return p


def _set_borders(table, color=BORDER, sz=4):
    tbl = table._tbl
    tblPr = tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), str(sz))
        e.set(qn("w:color"), color)
        borders.append(e)
    tblPr.append(borders)


def _no_borders(table):
    tblPr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "none")
        borders.append(e)
    tblPr.append(borders)


def _para(doc, text, size=11, bold=False, color=None, align=WD_ALIGN_PARAGRAPH.JUSTIFY,
          before=0, after=8):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    r = p.add_run(text)
    _set_run(r, size=size, bold=bold, color=color)
    return p


def _fix_styles(doc):
    """Noto Sans en docDefaults, Normal y todos los estilos; cuerpo 11 pt."""
    styles_el = doc.styles.element
    # docDefaults
    rPrDefault = styles_el.find(qn("w:docDefaults") + "/" + qn("w:rPrDefault"))
    dd = styles_el.find(qn("w:docDefaults"))
    if dd is not None:
        rprd = dd.find(qn("w:rPrDefault"))
        if rprd is not None:
            rpr = rprd.find(qn("w:rPr"))
            if rpr is not None:
                rf = rpr.find(qn("w:rFonts"))
                if rf is None:
                    rf = OxmlElement("w:rFonts")
                    rpr.insert(0, rf)
                for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
                    rf.set(qn(a), FONT)
                for tag, val in (("w:sz", "22"), ("w:szCs", "22")):
                    e = rpr.find(qn(tag))
                    if e is None:
                        e = OxmlElement(tag)
                        rpr.append(e)
                    e.set(qn("w:val"), val)
    # todos los estilos: familia Noto Sans (tamaños intactos); Normal = 11 pt
    for st in styles_el.findall(qn("w:style")):
        rpr = st.find(qn("w:rPr"))
        if rpr is None:
            continue
        rf = rpr.find(qn("w:rFonts"))
        if rf is None:
            rf = OxmlElement("w:rFonts")
            rpr.insert(0, rf)
        for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
            rf.set(qn(a), FONT)
        sid = st.get(qn("w:styleId"))
        if sid == "Normal":
            for tag in ("w:sz", "w:szCs"):
                e = rpr.find(qn(tag))
                if e is None:
                    e = OxmlElement(tag)
                    rpr.append(e)
                e.set(qn("w:val"), "22")


def _fix_fonttable_and_theme(path):
    """fontTable + theme latin -> Noto Sans (sobre el .docx ya guardado)."""
    import re
    tmp = path.with_suffix(".tmp.docx")
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.namelist():
            data = zin.read(item)
            if item == "word/fontTable.xml":
                xml = data.decode("utf-8")
                if "Noto Sans" not in xml:
                    entry = ('<w:font w:name="Noto Sans"><w:panose1 w:val="020B0604020202020204"/>'
                             '<w:charset w:val="00"/><w:family w:val="swiss"/>'
                             '<w:pitch w:val="variable"/><w:sig w:usb0="00000003" w:usb1="00000000" '
                             'w:usb2="00000000" w:usb3="00000000" w:csb0="0000009F" w:csb1="00000000"/></w:font>')
                    xml = xml.replace("</w:fonts>", entry + "</w:fonts>")
                data = xml.encode("utf-8")
            elif item == "word/theme/theme1.xml":
                xml = data.decode("utf-8")
                xml = re.sub(r'(<a:latin[^>]*typeface=")[^"]*(")', r"\1Noto Sans\2", xml)
                data = xml.encode("utf-8")
            elif item == "word/settings.xml":
                xml = data.decode("utf-8")
                # modo compatibilidad moderno (Word 2013+): evita "Modo de compatibilidad"
                if 'w:name="compatibilityMode"' not in xml:
                    xml = xml.replace(
                        "<w:compat>",
                        '<w:compat><w:compatSetting w:name="compatibilityMode" '
                        'w:uri="http://schemas.microsoft.com/office/word" w:val="15"/>')
                xml = xml.replace('<w:compatSetting w:name="compatibilityMode" '
                                  'w:uri="http://schemas.microsoft.com/office/word" w:val="14"/>',
                                  '<w:compatSetting w:name="compatibilityMode" '
                                  'w:uri="http://schemas.microsoft.com/office/word" w:val="15"/>')
                data = xml.encode("utf-8")
            elif item == "docProps/app.xml":
                xml = data.decode("utf-8")
                # metadato de páginas: la nota es exactamente 2 páginas
                xml = re.sub(r"<Pages>\d+</Pages>", "<Pages>2</Pages>", xml)
                data = xml.encode("utf-8")
            zout.writestr(item, data)
    tmp.replace(path)


def build(dest: Path):
    shutil.copy(PLANTILLA, dest)
    doc = Document(dest)
    _fix_styles(doc)

    # limpiar párrafos placeholder del cuerpo (se conserva sectPr)
    body = doc.element.body
    for p in list(body.findall(qn("w:p"))):
        body.remove(p)
    for t in list(body.findall(qn("w:tbl"))):
        body.remove(t)

    # ---- Banda 1: "Boletín de Indicador..." (barra delgada #10312B) ----
    t1 = doc.add_table(rows=1, cols=1)
    t1.alignment = WD_TABLE_ALIGNMENT.LEFT
    t1.autofit = False
    t1.columns[0].width = Twips(5240)
    t1.rows[0].cells[0].width = Twips(5240)
    _no_borders(t1)
    c = t1.rows[0].cells[0]
    _shade_cell(c, GREEN_DARK)
    _cell_para(c, [("Boletín de Indicador del 30 de julio de 2026  |  INEGI", 9, True, "FFFFFF")])

    # ---- Banda 2: título sobre banda institucional #235B4C ----
    t2 = doc.add_table(rows=1, cols=1)
    t2.alignment = WD_TABLE_ALIGNMENT.LEFT
    t2.autofit = False
    t2.columns[0].width = CONTENT_W
    t2.rows[0].cells[0].width = CONTENT_W
    _no_borders(t2)
    c = t2.rows[0].cells[0]
    _shade_cell(c, GREEN)
    _cell_para(c, [("Estimación Oportuna del PIB Trimestral", 15, True, "FFFFFF"),
                   ("PIB Oportuno · 2T 2026", 10.5, False, "FFFFFF")])

    # ---- KPIs: dos celdas claras ----
    tk = doc.add_table(rows=1, cols=2)
    tk.alignment = WD_TABLE_ALIGNMENT.LEFT
    tk.autofit = False
    for w, cell in zip((Twips(4150), Twips(4688)), tk.rows[0].cells):
        cell.width = w
    tk.columns[0].width = Twips(4150)
    tk.columns[1].width = Twips(4688)
    _no_borders(tk)
    for cell, (lbl, val) in zip(tk.rows[0].cells,
                                (("VARIACIÓN TRIMESTRAL", "▲ +1.5%"),
                                 ("VARIACIÓN ANUAL", "▲ +2.1%"))):
        _shade_cell(cell, CELL_BG)
        _cell_para(cell, [(lbl, 8.5, True, GREEN_DARK), (val, 15, True, GREEN_DARK)],
                   space_after=2)

    # ---- Tabla de actividades ----
    ta = doc.add_table(rows=3, cols=6)
    ta.alignment = WD_TABLE_ALIGNMENT.LEFT
    ta.autofit = False
    for col in ta.columns:
        col.width = Twips(1473)
    for cell in [c for r in ta.rows for c in r.cells]:
        cell.width = Twips(1473)
    _set_borders(ta)
    heads = ("Actividades primarias", "Actividades secundarias", "Actividades terciarias")
    for i in range(3):
        a = ta.rows[0].cells[i * 2]
        b = ta.rows[0].cells[i * 2 + 1]
        m = a.merge(b)
        _shade_cell(m, GREEN)
        _cell_para(m, [(heads[i], 9.5, True, "FFFFFF")])
    for j in range(6):
        c = ta.rows[1].cells[j]
        _shade_cell(c, CELL_BG)
        _cell_para(c, [("Trimestral" if j % 2 == 0 else "Anual", 8.5, False, GREEN_DARK)])
    vals = ("▲ 3.3%", "▲ 7.3%", "▲ 1.6%", "▲ 0.8%", "▲ 1.5%", "▲ 2.5%")
    for j, v in enumerate(vals):
        _cell_para(ta.rows[2].cells[j], [(v, 11, True, GREEN_DARK)])

    # ---- Texto introductorio ----
    _para(doc, "La estimación oportuna del PIB correspondiente al segundo trimestre de 2026 "
               "confirma que la economía mexicana recuperó dinamismo tras la contracción "
               "registrada a principios de año.", before=10, after=6)
    _para(doc, "Con cifras ajustadas por estacionalidad, el PIB creció 1.5% respecto al "
               "trimestre previo y 2.1% en su comparación anual, con lo que no sólo compensó "
               "la caída observada en el primer trimestre, sino que se ubicó por encima del "
               "nivel alcanzado al cierre de 2025. En conjunto, los datos son consistentes "
               "con los indicadores de coyuntura que anticipaban una recuperación de la "
               "actividad económica.", after=10)

    # ---- Gráfica 1 ----
    pimg = doc.add_paragraph()
    pimg.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pimg.paragraph_format.space_before = Pt(4)
    pimg.paragraph_format.space_after = Pt(0)
    pimg.add_run().add_picture(str(ASSETS / "chart_trimestral.png"), width=Cm(15.4))

    # ---- Salto a página 2 ----
    doc.add_page_break()

    # ---- Gráfica 2 (lleva título propio) ----
    pimg2 = doc.add_paragraph()
    pimg2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pimg2.paragraph_format.space_before = Pt(6)
    pimg2.paragraph_format.space_after = Pt(10)
    pimg2.add_run().add_picture(str(ASSETS / "chart_actividades.png"), width=Cm(15.4))

    # ---- Bloques de análisis ----
    _para(doc, "La recuperación fue de carácter generalizado. Las actividades primarias "
               "crecieron 3.3% trimestral, las secundarias 1.6% y las terciarias 1.5%, lo que "
               "muestra que el repunte no se concentró en un solo sector. Sin embargo, la "
               "comparación anual continúa evidenciando diferencias importantes: mientras las "
               "actividades primarias avanzaron 7.3% y las terciarias 2.5%, las secundarias "
               "registraron un crecimiento de apenas 0.8%, reflejando que la recuperación "
               "industrial aún es incompleta.", before=6, after=8)
    _para(doc, "En este sentido, las actividades secundarias continúan siendo el principal "
               "foco de debilidad de la economía mexicana, aunque el diagnóstico es más "
               "favorable que el observado hace tres meses. El crecimiento trimestral de la "
               "industria sugiere que el proceso de recuperación ya está en marcha; no "
               "obstante, su desempeño acumulado en el primer semestre permanece ligeramente "
               "por debajo del registrado un año antes, lo que indica que el sector todavía "
               "no recupera plenamente el dinamismo del resto de la economía.", after=8)
    _para(doc, "En conjunto, los resultados fortalecen el escenario de una economía que "
               "mantiene una trayectoria de crecimiento moderado y que logró superar la "
               "debilidad transitoria observada a principios de año. Hacia adelante, el "
               "principal elemento a seguir será la consolidación de la recuperación "
               "industrial, cuya evolución será determinante para sostener un crecimiento "
               "más equilibrado durante la segunda mitad de 2026.", after=0)

    doc.save(dest)
    _fix_fonttable_and_theme(dest)


def main():
    for dest in (MACHOTE, NOTA):
        dest.parent.mkdir(parents=True, exist_ok=True)
        build(dest)
        # verificación rápida
        with zipfile.ZipFile(dest) as z:
            assert "word/document.xml" in z.namelist()
            assert "word/settings.xml" in z.namelist()
        print("OK", dest.relative_to(ROOT))


if __name__ == "__main__":
    sys.exit(main())
