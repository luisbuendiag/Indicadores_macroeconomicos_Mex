"""Word de dos páginas con cuadros, bullets y visuales de encuestas."""
from docx import Document
from docx.shared import Pt,Inches,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH as ALIGN
import notas_word as W
from notas_imoam import period_text

def tables(f):
    rows=f['rows']
    if f['indicator']=='EMIM':
        selected=sorted(f['sectors'],key=lambda r:(r['production'],r['code']),reverse=True)
        selected=selected[:3]+selected[-3:]
        return [
            [['Variable','Mensual (%)','Anual (%)']]+[[r['label'],f'{r["mom"]:.1f}',f'{r["yoy"]:.1f}'] for r in rows],
            [['Variable','Índice original','Anual original (%)']]+[[r['label'],f'{r["index"]:.1f}',f'{r["original"]:.1f}'] for r in rows],
            [['Subsector','Producción (%)','Empleo (%)']]+[[r['label'],f'{r["production"]:.1f}',f'{r["employment"]:.1f}'] for r in selected]]
    def c(sector,label):return next(r['level'] for r in f['components'] if r['sector']==sector and r['label']==label)
    return [
        [['Sector y serie','Nivel (pts)','Mensual (pts)','Anual (pts)']]+[[r['label']+(' (D)' if r['series']=='desestacionalizada' else ' (O)'),f'{r["level"]:.1f}',f'{r["mom"]:.1f}',f'{r["yoy"]:.1f}'] for r in rows],
        [['Sector','Invertir¹','Empresa hoy','Empresa futura']]+[[r['label']]+[f'{c(r["label"],s):.1f}' for s in ['Momento para invertir','Situación presente de la empresa','Situación futura de la empresa']] for r in rows]]

def build(f,url,template,out):
    W.fonts()
    import matplotlib.pyplot as plt
    doc=Document(template)
    for node in list(doc.element.body):
        if node.tag!=qn('w:sectPr'):doc.element.body.remove(node)
    for st in doc.styles:
        if st.type==1:st.font.name='Noto Sans';st.font.size=Pt(11)
    def p(text='',size=11,bold=False,color='000000',keep=False):
        x=doc.add_paragraph();x.alignment=ALIGN.JUSTIFY
        pf=x.paragraph_format;pf.space_before=Pt(0);pf.space_after=Pt(0);pf.line_spacing=1;pf.keep_with_next=keep
        r=x.add_run(text);r.font.name='Noto Sans';r.font.size=Pt(size);r.bold=bold;r.font.color.rgb=RGBColor.from_string(color);return x
    def h(t):p('',6.5);p(t,bold=True,color='235B4C',keep=True)
    def note(t):p(t,8,color='666666')
    # Lista nativa independiente de estilos/IDs específicos del machote.
    numbering=doc.part.numbering_part.element
    aid=max([int(n.get(qn('w:abstractNumId'))) for n in numbering.findall(qn('w:abstractNum'))]+[-1])+1
    nid=max([int(n.get(qn('w:numId'))) for n in numbering.findall(qn('w:num'))]+[0])+1
    a=OxmlElement('w:abstractNum');a.set(qn('w:abstractNumId'),str(aid));lvl=OxmlElement('w:lvl');lvl.set(qn('w:ilvl'),'0')
    for tag,val in [('numFmt','bullet'),('lvlText','•'),('suff','space')]:
        n=OxmlElement('w:'+tag);n.set(qn('w:val'),val);lvl.append(n)
    a.append(lvl);numbering.append(a);n=OxmlElement('w:num');n.set(qn('w:numId'),str(nid));x=OxmlElement('w:abstractNumId');x.set(qn('w:val'),str(aid));n.append(x);numbering.append(n)
    def bullet(t):
        x=p(t);pr=OxmlElement('w:numPr')
        for tag,val in [('ilvl','0'),('numId',str(nid))]:
            n=OxmlElement('w:'+tag);n.set(qn('w:val'),val);pr.append(n)
        x._p.get_or_add_pPr().append(pr)
    def table(data,hero=False):
        t=doc.add_table(rows=0,cols=len(data[0]));t.autofit=False
        widths={2:[3.05,3.05],3:[3.1,1.5,1.5],4:[2.65,1.15,1.15,1.15]}[len(data[0])]
        for col,w in zip(t.columns,widths):col.width=Inches(w)
        for i,row in enumerate(data):
            rr=t.add_row();rr._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
            for c,txt,w in zip(rr.cells,row,widths):
                c.width=Inches(w);c.text=str(txt);shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'10312B' if hero or i==0 else ('F4F4F2' if i%2 else 'FFFFFF'));c._tc.get_or_add_tcPr().append(shade)
                for x in c.paragraphs:
                    x.paragraph_format.space_before=Pt(0);x.paragraph_format.space_after=Pt(0);x.paragraph_format.line_spacing=1
                    for r in x.runs:r.font.name='Noto Sans';r.font.size=Pt(10);r.bold=hero or i==0;r.font.color.rgb=RGBColor.from_string('FFFFFF' if hero or i==0 else '000000')
        return t
    def plot(name,labels,vals,title,unit,other=None):
        with plt.rc_context({'font.family':'Noto Sans','font.size':9}):
            fig,ax=plt.subplots(figsize=(6.1,1.75 if len(vals)<=4 else 1.9));ys=list(range(len(vals)))
            if other is None:ax.barh(ys,vals,color=['#235B4C' if v>=0 else '#611232' for v in vals])
            else:
                ax.barh([y-.17 for y in ys],vals,height=.32,color='#235B4C',label='Empresa presente')
                ax.barh([y+.17 for y in ys],other,height=.32,color='#A57F2C',label='Empresa futura');ax.axvline(50,color='#611232',ls='--',lw=1)
            ax.set_yticks(ys,labels);ax.invert_yaxis();ax.set_title(title,fontsize=11,color='#10312B');ax.set_xlabel(unit)
            ax.axvline(0,color='#888888',lw=.6)
            values=vals+(other or []);pad=max(1,max(abs(v) for v in values))*.25;ax.set_xlim(min(0,min(values))-pad,max(0,max(values))+pad)
            for i,v in enumerate(vals):ax.text(v,i if other is None else i-.17,f' {v:.1f}' if v>=0 else f'{v:.1f} ',va='center',ha='left' if v>=0 else 'right',fontsize=8)
            if other:
                for i,v in enumerate(other):ax.text(v,i+.17,f' {v:.1f}',va='center',fontsize=8)
            ax.spines[['top','right']].set_visible(False);fig.tight_layout();path=out.parent/(name+'.png');fig.savefig(path,dpi=300);plt.close(fig)
        doc.add_picture(str(path),width=Inches(6.05))
    key=f['indicator'];rows=f['rows'];data=tables(f)
    hero=rows[0] if key=='EMIM' else f['globalrow']
    table([[key+' · '+period_text(f['ym']).capitalize(),'Publicación '+f['publication_date']],
           [f'Producción mensual {hero["mom"]:.1f}%' if key=='EMIM' else f'IGOEC {hero["level"]:.1f} puntos',f'Anual {hero["yoy"]:.1f}'+('%' if key=='EMIM' else ' puntos')]],True)
    p('Encuesta Mensual de la Industria Manufacturera' if key=='EMIM' else 'Confianza empresarial',bold=True,color='10312B')
    h('Resumen ejecutivo')
    if key=='EMIM':
        p(f'En {period_text(f["ym"])}, la producción manufacturera varió {rows[0]["mom"]:.1f}% mensual y {rows[0]["yoy"]:.1f}% anual, con cifras desestacionalizadas.')
        bullet(f'Las remuneraciones medias reales registraron {rows[3]["yoy"]:.1f}% anual; en el mes variaron {rows[3]["mom"]:.1f}%.')
        bullet(f'El personal ocupado varió {rows[1]["mom"]:.1f}% mensual y {rows[1]["yoy"]:.1f}% anual; las horas trabajadas, {rows[2]["mom"]:.1f}% y {rows[2]["yoy"]:.1f}%, respectivamente.')
        h('Producción y condiciones laborales')
        table(data[0]);note('Fuente: INEGI, cuadro 1, p. 3. Comparaciones desestacionalizadas; tasas publicadas, sin sumar componentes.')
        plot('anual',[r['label'] for r in rows],[r['yoy'] for r in rows],'Balance anual de las cuatro variables','Variación anual (%) · cifras desestacionalizadas')
        h('Cifras originales del sector manufacturero');table(data[1]);note('Fuente: cuadro 2, p. 4. Índice '+f['base']+'=100; cifras preliminares. No son los índices desestacionalizados de la gráfica.')
        doc.add_page_break();h('Desempeño por subsector')
        sectors=f['sectors'];positive=sum(r['production']>0 for r in sectors);both=sum(r['production']>0 and r['employment']>0 for r in sectors)
        p(f'{positive} de {len(sectors)} subsectores aumentaron su producción frente al mismo mes del año anterior; {both} registraron simultáneamente crecimiento de producción y empleo. Conteos no ponderados, en cifras originales.')
        table(data[2]);note('Fuente: cuadro 2, p. 4. Tres mayores y tres menores tasas anuales de producción entre los 21 subsectores; empleo para la misma selección.')
        selected=sorted(sectors,key=lambda r:(r['production'],r['code']),reverse=True);selected=selected[:3]+selected[-3:]
        plot('subsectores',[r['label'] for r in selected],[r['production'] for r in selected],'Contrastes en la producción manufacturera','Variación anual (%) · cifras originales')
        h('Lectura sectorial')
        best=selected[0];worst=selected[-1]
        bullet(f'La mayor tasa anual de producción fue la de {best["label"].lower()} ({best["production"]:.1f}%); su empleo varió {best["employment"]:.1f}%. Una tasa elevada no identifica la mayor aportación al total.')
        bullet(f'La menor tasa de producción fue la de {worst["label"].lower()} ({worst["production"]:.1f}%), con {worst["employment"]:.1f}% en personal ocupado.')
        h('Notas metodológicas')
        p('Las remuneraciones medias reales están deflactadas con el INPC; no representan la masa salarial. Producción, personal y horas tienen distinta cobertura conceptual: no se deriva productividad con estos agregados.')
        note('Cifras originales preliminares. Los agregados desestacionalizados se ajustan de forma independiente. El empleo EMIM no equivale a puestos afiliados al IMSS. La encuesta manufacturera no mide toda la actividad industrial.')
    else:
        p(f'En {period_text(f["ym"])}, el IGOEC registró {hero["level"]:.1f} puntos, con una diferencia mensual de {hero["mom"]:.1f} puntos y anual de {hero["yoy"]:.1f} puntos, en cifras desestacionalizadas.')
        positive=[r['label'].lower() for r in rows if r['mom']>0];negative=[r['label'].lower() for r in rows if r['mom']<0]
        gains=', '.join(positive) if positive else 'ninguno'
        losses=', '.join(negative) if negative else 'ninguno'
        bullet(f'Sectores con aumento mensual: {gains}. Sectores con disminución: {losses}.')
        below=sum(r['level']<50 for r in rows)
        bullet(f'{below} de cuatro indicadores sectoriales se ubicaron por debajo del umbral de 50 puntos. Las opiniones describen percepciones empresariales, no producción observada.')
        h('Confianza por sector');table(data[0]);note('D: desestacionalizada. O: original, sin patrón estacional definido en esta edición. Fuente: cuadros 1–4, pp. 2–4. Las diferencias están en puntos, no en porcentajes.')
        plot('mensual',[r['label'] for r in rows],[r['mom'] for r in rows],'Diferencia mensual de la confianza por sector','Puntos · D: Manufacturas; O: los otros tres sectores')
        h('Lectura del umbral')
        p('El ICE se centra en 50 puntos. Valores superiores reflejan una percepción favorable; valores inferiores, una expectativa menos favorable. El IGOEC se calcula con ponderaciones sectoriales y no es el promedio simple de los niveles de este cuadro.')
        doc.add_page_break();h('Percepciones de inversión y de la empresa')
        table(data[1]);note('Niveles en puntos. ¹ Inversión: manufacturas y construcción ajustadas; comercio y servicios originales. Empresa presente y futura: originales en los cuatro sectores. Fuente: cuadros 1–4.')
        def comp(label):return [next(c['level'] for c in f['components'] if c['sector']==r['label'] and c['label']==label) for r in rows]
        present=comp('Situación presente de la empresa');future=comp('Situación futura de la empresa')
        plot('expectativas',[r['label'] for r in rows],present,'Situación presente y futura de la empresa','Puntos · presente verde y futura dorado · originales · umbral 50',future)
        h('Balance de expectativas')
        bullet(f'La situación futura de la empresa se ubicó sobre 50 puntos en {sum(v>50 for v in future)} de cuatro sectores; la presente, en {sum(v>50 for v in present)}. Son respuestas a preguntas distintas, no una previsión de crecimiento.')
        invest=comp('Momento para invertir');bullet(f'El componente de momento adecuado para invertir estuvo bajo 50 puntos en {sum(v<50 for v in invest)} de cuatro sectores. El nivel expresa la opinión sobre invertir, no una medición de inversión realizada.')
        h('Notas metodológicas')
        p('El módulo corresponde al boletín de confianza derivado de la EMOE. No mezcla los indicadores de pedidos ni de tendencia. Las diferencias publicadas se calculan con todos los decimales de los indicadores; pueden diferir de restar niveles redondeados.')
        p('En esta edición, el ICE manufacturero y el IGOEC se presentan ajustados; construcción, comercio y servicios utilizan cifras originales. La clasificación de cada componente se conserva por separado. Los resultados originales del mes son preliminares.')
        note('La situación presente de la empresa es una percepción del informante; la situación futura es una expectativa. No se atribuyen causalidad, crecimiento observado ni aportaciones sectoriales a estos índices.')
    p('',6.5)
    note(f'INEGI, boletín {f["bulletin"]}, publicado el {f["publication_date"]}. Cuadros y metodología de la misma edición. Cifras sujetas a revisión.')
    note(url);doc.core_properties.title=key+' '+f['period'];doc.save(out)

def validate(f,doc):
    expected=tables(f)
    if len(doc.tables)!=len(expected)+1 or len(doc.inline_shapes)!=2:raise ValueError('Faltan cuadros o visuales')
    for actual,wanted in zip(doc.tables[1:],expected):
        if [[c.text for c in r.cells] for r in actual.rows]!=wanted:raise ValueError('El cuadro Word difiere del boletín')
