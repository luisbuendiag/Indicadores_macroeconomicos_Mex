"""Word nativo y render del piloto: plantilla de solo lectura y QA por página."""
from __future__ import annotations
import os
import shutil
import subprocess
import tempfile
from pathlib import Path

import pdfplumber
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH as ALIGN
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from notas_imoam import number, period_text



def fonts():
    import matplotlib
    matplotlib.use('Agg')
    from matplotlib import font_manager
    font=os.environ.get('NOTES_FONT_FILE')
    if font:font_manager.fontManager.addfont(font)
    font_manager.findfont('Noto Sans',fallback_to_default=False)


def interval_plot(facts: dict,out: Path) -> None:
    fonts()
    import matplotlib.pyplot as plt
    with plt.rc_context({'font.family':'Noto Sans','font.size':9}):
        fig,ax=plt.subplots(figsize=(6.1,1.65))
        lo,hi,center=[facts['values'][i]*100 for i in (4,5,3)]
        ax.errorbar(center,0,xerr=[[center-lo],[hi-center]],fmt='o',color='#10312B',ecolor='#235B4C',capsize=8,linewidth=2,markersize=7)
        ax.axvline(0,color='#999999',linestyle='--',linewidth=1)
        ax.annotate(f'Estimación: {center:.1f}%',(center,0),xytext=(0,15),textcoords='offset points',ha='center',fontweight='bold',color='#10312B')
        pad=max((hi-lo)*.13,.3);ax.set_xlim(min(lo,0)-pad,max(hi,0)+pad);ax.set_ylim(-.3,.45)
        ax.set_yticks([]);ax.set_xlabel('Variación anual (%) · cifras originales · IC al 95%')
        ticks=sorted(set([lo,0,hi]));ax.set_xticks(ticks);ax.set_xticklabels([f'{x:.1f}%' for x in ticks])
        ax.set_title('IMOAM · '+period_text(facts['ym']).capitalize(),fontsize=11,color='#10312B',fontweight='bold',pad=12)
        for spine in ax.spines.values():spine.set_visible(False)
        fig.tight_layout();fig.savefig(out,dpi=300,facecolor='white');plt.close(fig)


def opportunity_plot(facts: dict,out: Path) -> None:
    fonts()
    import matplotlib.pyplot as plt
    from matplotlib.patches import FancyBboxPatch
    c=facts['context']
    with plt.rc_context({'font.family':'Noto Sans','font.size':9}):
        fig,ax=plt.subplots(figsize=(6.1,1.55));ax.set_xlim(0,1);ax.set_ylim(0,1);ax.axis('off')
        boxes=[(.01,'Mes de referencia','Cierre del mes'),(.35,'IMOAM','Estimación oportuna'),(.69,'IMAI manufacturero',f'Aprox. {c["imai_lag"]} días tras cierre')]
        for x,title,sub in boxes:
            ax.add_patch(FancyBboxPatch((x,.32),.29,.43,boxstyle='round,pad=0.007',edgecolor='#235B4C',facecolor='#F4F4F2'))
            ax.text(x+.145,.60,title,ha='center',va='center',fontsize=9.5,fontweight='bold',color='#10312B')
            ax.text(x+.145,.43,sub,ha='center',va='center',fontsize=8.3,color='#4C545B')
        for x in (.31,.65):ax.annotate('',xy=(x+.035,.54),xytext=(x,.54),arrowprops={'arrowstyle':'->','color':'#235B4C','lw':1.5})
        ax.text(.66,.15,f'El IMOAM anticipa al IMAI en aproximadamente {c["days_ahead"]} días',ha='center',fontsize=9,color='#10312B')
        fig.tight_layout(pad=.5);fig.savefig(out,dpi=300,facecolor='white');plt.close(fig)


def build(facts: dict,url: str,template: Path,out: Path) -> str:
    from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
    from docx.oxml import OxmlElement
    from sources.inegi_bulletin import MESES_NOMBRE
    if not facts.get('context'):raise ValueError('Faltan los datos ampliados de evaluación y metodología')
    v,c=facts['values'],facts['context']
    interval_plot(facts,out.parent/'intervalo.png');opportunity_plot(facts,out.parent/'oportunidad.png')
    doc=Document(template)
    for child in list(doc._element.body):
        if child.tag!=qn('w:sectPr'):doc._element.body.remove(child)
    for name in ('Normal','Title','Heading 1','List Bullet'):
        if name not in doc.styles:
            from docx.enum.style import WD_STYLE_TYPE
            doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        st=doc.styles[name];st.font.name='Noto Sans';st.font.size=Pt(11)
        st.paragraph_format.space_before=Pt(0);st.paragraph_format.space_after=Pt(0);st.paragraph_format.line_spacing=1
    settings=doc.settings.element
    for n in settings.findall('.//'+qn('w:useFELayout')):n.getparent().remove(n)
    compat=settings.find(qn('w:compat'))
    if compat is None:compat=OxmlElement('w:compat');settings.append(compat)
    for n in list(compat):
        if n.get(qn('w:name'))=='compatibilityMode':compat.remove(n)
    mode=OxmlElement('w:compatSetting')
    for k,val in [('name','compatibilityMode'),('uri','http://schemas.microsoft.com/office/word'),('val','15')]:mode.set(qn('w:'+k),val)
    compat.append(mode)
    numbering=doc.part.numbering_part.element
    aid=max([int(n.get(qn('w:abstractNumId'))) for n in numbering.findall(qn('w:abstractNum'))]+[-1])+1
    nid=max([int(n.get(qn('w:numId'))) for n in numbering.findall(qn('w:num'))]+[0])+1
    abstract=OxmlElement('w:abstractNum');abstract.set(qn('w:abstractNumId'),str(aid))
    level=OxmlElement('w:lvl');level.set(qn('w:ilvl'),'0')
    for tag,attrs in [('start',{'val':'1'}),('numFmt',{'val':'bullet'}),('lvlText',{'val':'•'}),('lvlJc',{'val':'left'}),('suff',{'val':'space'})]:
        n=OxmlElement('w:'+tag)
        for k,value in attrs.items():n.set(qn('w:'+k),value)
        level.append(n)
    abstract.append(level);numbering.append(abstract)
    num=OxmlElement('w:num');num.set(qn('w:numId'),str(nid));ref=OxmlElement('w:abstractNumId');ref.set(qn('w:val'),str(aid));num.append(ref);numbering.append(num)
    audit=[]
    def run(p,text,size=11,bold=False,color='000000',italic=False):
        r=p.add_run(text);r.font.name='Noto Sans';r.font.size=Pt(size);r.bold=bold;r.italic=italic;r.font.color.rgb=RGBColor.from_string(color);return r
    def pf(p,align=ALIGN.JUSTIFY,keep=False):
        p.alignment=align;f=p.paragraph_format;f.space_before=Pt(0);f.space_after=Pt(0);f.line_spacing=1;f.keep_with_next=keep;f.widow_control=True
    def para(text='',size=11,bold=False,color='000000',style=None,keep=False):
        p=doc.add_paragraph(style=style);pf(p,keep=keep);run(p,text,size,bold,color);audit.append(text);return p
    def gap():para('',6.5)
    def heading(text):gap();para(text,11,True,'235B4C',keep=True)
    def note(text):
        p=para(text,8,color='666666');p.runs[0].italic=True
    def bullet(text):
        p=para(text,style='List Bullet')
        np=OxmlElement('w:numPr');il=OxmlElement('w:ilvl');il.set(qn('w:val'),'0');ni=OxmlElement('w:numId');ni.set(qn('w:val'),str(nid));np.extend([il,ni]);p._p.get_or_add_pPr().append(np)
        p.paragraph_format.left_indent=Inches(.14);p.paragraph_format.first_line_indent=Inches(-.10)
    def shade(cell,fill):
        pr=cell._tc.get_or_add_tcPr();n=OxmlElement('w:shd');n.set(qn('w:fill'),fill);pr.append(n)
        cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    def cell_text(cell,lines,fill,align=ALIGN.CENTER):
        cell.text='';shade(cell,fill)
        for i,(text,size,bold,color) in enumerate(lines):
            p=cell.paragraphs[0] if i==0 else cell.add_paragraph();pf(p,align)
            run(p,text,size,bold,color);audit.append(text)
    def table(rows,widths,header=True):
        t=doc.add_table(rows=0,cols=len(widths));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
        for col,w in zip(t.columns,widths):col.width=Inches(w)
        for ri,row in enumerate(rows):
            cells=t.add_row().cells
            for ci,(cell,text,w) in enumerate(zip(cells,row,widths)):
                cell.width=Inches(w)
                fill='10312B' if ri==0 and header else ('F4F4F2' if ri%2 else 'FFFFFF')
                color='FFFFFF' if ri==0 and header else '000000'
                cell_text(cell,[(text,10.5,ri==0 or ci==0,color)],fill,ALIGN.LEFT if ci==0 else ALIGN.CENTER)
            pr=t.rows[-1]._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
            if ri==0 and header:pr.append(OxmlElement('w:tblHeader'))
        return t
    def figure(path):
        p=doc.add_paragraph();pf(p,ALIGN.CENTER);p.add_run().add_picture(str(path),width=Inches(6.05))
    # Cabecera visual derivada del Word ENOE aportado por el usuario.
    year,month,day=map(int,facts['publication_date'].split('-'))
    publication=f'{day} de {MESES_NOMBRE[month-1]} de {year}'
    hero=doc.add_table(rows=3,cols=2);hero.alignment=WD_TABLE_ALIGNMENT.CENTER;hero.autofit=False
    for col in hero.columns:col.width=Inches(3.065)
    cell=hero.cell(0,0).merge(hero.cell(0,1))
    cell_text(cell,[(f'Boletín del {publication} | INEGI',10,True,'FFFFFF')],'10312B')
    cell=hero.cell(1,0).merge(hero.cell(1,1))
    cell_text(cell,[('Indicador Mensual Oportuno de la Actividad Manufacturera',11,False,'FFFFFF'),('IMOAM · '+period_text(facts['ym']).replace(' de ',' ').capitalize(),20,True,'FFFFFF')],'235B4C')
    for cell,label,value,sub in [(hero.cell(2,0),'ÍNDICE ESTIMADO',number(v[0]),'Puntos · cifras originales'),(hero.cell(2,1),'VARIACIÓN ANUAL ESTIMADA',number(v[3],True),'Respecto al mismo mes del año anterior')]:
        cell_text(cell,[(label,9.5,True,'FFFFFF'),(value,26,True,'FFFFFF'),(sub,8.5,False,'FFFFFF')],'538135')
    for row in hero.rows:row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
    note('Nota: estadística experimental. El IMOAM anticipa el IMAI manufacturero; no representa un dato observado definitivo.')
    heading('Resumen ejecutivo')
    para('La estimación puntual apunta a '+('un aumento' if v[3]>0 else 'una disminución' if v[3]<0 else 'estabilidad')+' de la actividad manufacturera en su comparación anual. La lectura debe considerar los intervalos de confianza y el carácter oportuno del indicador.')
    cross=v[4]<=0<=v[5]
    bullet(('El intervalo anual incluye cero y valores negativos: la señal puntual positiva no implica certeza de crecimiento.' if v[3]>0 and v[4]<0<=v[5] else 'El intervalo anual incluye cero: la estimación está sujeta a incertidumbre sobre el signo del cambio.') if cross else ('El intervalo anual permanece por encima de cero.' if v[4]>0 else 'El intervalo anual permanece por debajo de cero.'))
    bullet(f'La amplitud del intervalo anual es de {number((v[5]-v[4])*100)} puntos porcentuales; la del índice es de {number(v[2]-v[1])} puntos.')
    bullet(f'La publicación aporta una lectura anticipada: se difunde alrededor de {c["days_ahead"]} días antes del IMAI manufacturero, según la ficha metodológica.')
    heading('Estimación puntual e intervalos de confianza')
    table([['Variable','Estimación','Límite inferior','Límite superior'],['Índice (puntos)',number(v[0]),number(v[1]),number(v[2])],['Variación anual (%)',number(v[3],True),number(v[4],True),number(v[5],True)]],[2.35,1.26,1.26,1.26])
    note('IC al 95%. Fuente: INEGI, cuadro 1, página 1. Amplitud = límite superior − límite inferior; cálculo a partir de las cifras publicadas.')
    figure(out.parent/'intervalo.png')
    note('Fuente: INEGI, IMOAM. El punto indica la estimación; el segmento muestra el intervalo al 95%.')
    # Segunda parte: contenido sustantivo, sin salto de página impuesto.
    heading('Evaluación histórica y oportunidad de publicación')
    para(f'INEGI reporta que el IMAI manufacturero quedó dentro del intervalo de confianza del IMOAM en {c["coverage"]:g}% de las ocasiones. La comparación presentada en esta edición abarca de {period_text(c["evaluation_start"])} a {period_text(c["evaluation_end"])}.')
    table([['Dato de evaluación','Resultado reportado'],['Cobertura histórica del intervalo',f'{c["coverage"]:g}% de las ocasiones'],['Inicio del cálculo del IMOAM',period_text(c['calculation_start']).capitalize()],['Comparación de referencia','IMOAM frente al IMAI del sector manufacturero'],['Nivel nominal del intervalo','95% de confianza']],[2.65,3.48])
    note('Fuente: INEGI, evaluación histórica, página 3. La cobertura observada describe desempeño pasado; no es la probabilidad de que la estimación de este mes resulte correcta.')
    figure(out.parent/'oportunidad.png')
    note('Fuente: INEGI, ficha metodológica, página 4. Esquema de difusión sin escala temporal exacta; los plazos son aproximados.')
    heading('Cobertura y construcción del indicador')
    table([['Elemento','Alcance de la edición'],['Unidad de observación',c['observation_unit']],['Variable que se anticipa','IMAI del sector manufacturero'],['Insumos del modelo','RENEM y registros de consumo de energía eléctrica de CFE y CENACE'],['Tratamiento estadístico','Serie mensual original, sin ajuste estacional'],['Disponibilidad de insumos','La serie mensual del ICEE no se publica por confidencialidad']],[2.05,4.08])
    heading('Notas de interpretación')
    bullet('La variación anual compara el mes de referencia con el mismo mes del año anterior; no equivale a una variación mensual desestacionalizada.')
    bullet('El IMOAM es una estimación del IMAI manufacturero. No debe confundirse con la EMIM ni atribuirse a ramas o entidades sin un desglose oficial.')
    note('Fuentes: INEGI, boletín '+str(facts['bulletin'] or 'IMOAM')+', '+publication+'. Resultados e intervalos: p. 1; evaluación histórica: p. 3; metodología: p. 4. El modelo utiliza registros administrativos para ofrecer información oportuna; estos resultados no prueban una relación causal.')
    note(url)
    doc.core_properties.title='IMOAM · '+period_text(facts['ym'])
    doc.core_properties.subject='Nota ejecutiva de dos páginas; formato visual basado en la referencia ENOE'
    doc.save(out)
    return '\n\n'.join(audit)

def binary(name: str, env: str) -> str:
    path=os.environ.get(env) or shutil.which(name)
    if not path:raise RuntimeError(f'Falta herramienta de render: {name}')
    return path


def raster(pdf: Path, prefix: Path, dpi=120) -> list[Path]:
    subprocess.run([binary('pdftoppm','NOTES_PDFTOPPM'),'-r',str(dpi),'-png',str(pdf),str(prefix)],check=True,capture_output=True,timeout=120)
    pages=sorted(prefix.parent.glob(prefix.name+'-*.png'))
    if not pages:raise ValueError('Render sin imágenes')
    return pages


def render(docx: Path, max_pages=2) -> tuple[Path,list[Path]]:
    docx.with_suffix('.pdf').unlink(missing_ok=True)
    with tempfile.TemporaryDirectory(prefix='notes-lo-') as profile:
        subprocess.run([binary('soffice','NOTES_SOFFICE'),'-env:UserInstallation='+Path(profile).as_uri(),'--headless','--convert-to','pdf','--outdir',str(docx.parent),str(docx)],check=True,capture_output=True,timeout=120)
    pdf=docx.with_suffix('.pdf')
    if not pdf.exists():raise ValueError('No se generó el PDF de revisión')
    with pdfplumber.open(pdf) as d:
        if not 1<=len(d.pages)<=max_pages:raise ValueError('Extensión de la nota fuera del límite')
        for page in d.pages:
            if len(page.extract_text() or '')<100:raise ValueError('Página vacía o ilegible')
            if any(c['x0']<0 or c['x1']>page.width+1 or c['top']<0 or c['bottom']>page.height+1 for c in page.chars):
                raise ValueError('Texto fuera de la página')
        fonts={c['fontname'] for pg in d.pages for c in pg.chars if c['text'].strip()}
        if not any('NotoSans' in n.replace('-','') for n in fonts):
            raise ValueError('Render sin Noto Sans; revisar fuentes instaladas')
    return pdf,raster(pdf,docx.parent/'nota-pagina')
