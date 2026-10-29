"""Composición institucional de notas IMAI e IGAE con cifras del boletín."""
from pathlib import Path
from docx import Document
from docx.shared import Pt,Inches,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH as A
import notas_word as W
from notas_imoam import period_text

def build(f,url,template,out):
    W.fonts()
    import matplotlib.pyplot as plt
    rows=f['rows'];parts=rows[1:];key=f['indicator'];total=rows[0]
    labels=[r['label'].replace('Generación, transmisión','Energía, agua y gas').replace('Industrias manufactureras','Manufacturas') for r in parts]
    for field,title in [('mom','Variación mensual desestacionalizada'),('yoy','Variación anual desestacionalizada')]:
        with plt.rc_context({'font.family':'Noto Sans','font.size':9}):
            fig,ax=plt.subplots(figsize=(6.1,1.9));vals=[r[field] for r in parts]
            ax.barh(labels,vals,color=['#235B4C' if v>=0 else '#611232' for v in vals]);ax.axvline(0,color='#777777',lw=.7)
            ax.invert_yaxis();ax.set_xlabel('Variación real (%)');ax.set_title(title,fontsize=11,color='#10312B')
            pad=max(1,max(abs(v) for v in vals))*.2;ax.set_xlim(min(0,min(vals))-pad,max(0,max(vals))+pad)
            for i,v in enumerate(vals):ax.text(v,i,f' {v:.1f}%' if v>=0 else f'{v:.1f}% ',ha='left' if v>=0 else 'right',va='center',fontsize=9)
            ax.spines[['top','right']].set_visible(False);fig.tight_layout();fig.savefig(out.parent/(field+'.png'),dpi=300);plt.close(fig)
    doc=Document(template)
    for n in list(doc.element.body):
        if n.tag!=qn('w:sectPr'):doc.element.body.remove(n)
    for st in doc.styles:
        if st.type==1:st.font.name='Noto Sans';st.font.size=Pt(11)
    def p(text='',size=11,bold=False,color='000000',keep=False):
        x=doc.add_paragraph();x.alignment=A.JUSTIFY;x.paragraph_format.space_after=Pt(0);x.paragraph_format.space_before=Pt(0);x.paragraph_format.line_spacing=1;x.paragraph_format.keep_with_next=keep
        r=x.add_run(text);r.font.name='Noto Sans';r.font.size=Pt(size);r.bold=bold;r.font.color.rgb=RGBColor.from_string(color);return x
    def h(text):p('',6.5);p(text,11,True,'235B4C',True)
    def note(text):p(text,8,color='666666')
    def table(data,hero=False):
        t=doc.add_table(rows=0,cols=len(data[0]));t.autofit=False
        widths=[6.1/len(data[0])]*len(data[0])
        for c,w in zip(t.columns,widths):c.width=Inches(w)
        for i,row in enumerate(data):
            rr=t.add_row();rr._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
            for c,txt in zip(rr.cells,row):
                c.text=str(txt);pr=c._tc.get_or_add_tcPr();sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'10312B' if i==0 or hero else ('F4F4F2' if i%2 else 'FFFFFF'));pr.append(sh)
                for x in c.paragraphs:
                    x.paragraph_format.space_after=Pt(0);x.paragraph_format.space_before=Pt(0);x.paragraph_format.line_spacing=1
                    for r in x.runs:r.font.name='Noto Sans';r.font.size=Pt(10.5);r.bold=i==0 or hero;r.font.color.rgb=RGBColor.from_string('FFFFFF' if i==0 or hero else '000000')
        return t
    numbering=doc.part.numbering_part.element
    aid=max([int(n.get(qn('w:abstractNumId'))) for n in numbering.findall(qn('w:abstractNum'))]+[-1])+1
    nid=max([int(n.get(qn('w:numId'))) for n in numbering.findall(qn('w:num'))]+[0])+1
    abstract=OxmlElement('w:abstractNum');abstract.set(qn('w:abstractNumId'),str(aid));lvl=OxmlElement('w:lvl');lvl.set(qn('w:ilvl'),'0')
    for tag,value in [('numFmt','bullet'),('lvlText','•'),('suff','space')]:
        n=OxmlElement('w:'+tag);n.set(qn('w:val'),value);lvl.append(n)
    abstract.append(lvl);numbering.append(abstract);num=OxmlElement('w:num');num.set(qn('w:numId'),str(nid));n=OxmlElement('w:abstractNumId');n.set(qn('w:val'),str(aid));num.append(n);numbering.append(num)
    def bullet(text):
        x=p(text);np=OxmlElement('w:numPr');il=OxmlElement('w:ilvl');il.set(qn('w:val'),'0');ni=OxmlElement('w:numId');ni.set(qn('w:val'),str(nid));np.extend([il,ni]);x._p.get_or_add_pPr().append(np)
    title='Indicador Mensual de la Actividad Industrial' if key=='IMAI' else 'Indicador Global de la Actividad Económica'
    table([[key+' · '+period_text(f['ym']).capitalize(), 'Publicación: '+f['publication_date']],[f'Índice desestacionalizado: {total["index"]:.1f}',f'Mensual: {total["mom"]:.1f}% | Anual: {total["yoy"]:.1f}%']],True)
    p(title,11,True,'10312B');note('Índice '+f['base']+'=100. Cifras de la edición del boletín; pueden diferir de las series revisadas posteriormente.')
    h('Resumen ejecutivo')
    p(f'En {period_text(f["ym"])}, el {key} registró una variación real mensual de {total["mom"]:.1f}% y anual de {total["yoy"]:.1f}%, con cifras desestacionalizadas. El índice se ubicó en {total["index"]:.1f} puntos.')
    best=max(parts,key=lambda r:r['mom']);worst=min(parts,key=lambda r:r['yoy'])
    leaders=', '.join(r['label'].replace('Generación, transmisión','Energía, agua y gas').lower() for r in parts if r['mom']==best['mom'])
    bullet(f'Mayor variación mensual entre los componentes principales: {leaders}, con {best["mom"]:.1f}%.')
    bullet(f'Menor variación anual entre los componentes principales: {worst["label"].lower()}, con {worst["yoy"]:.1f}%. Esta tasa describe su variación, no su contribución al total.')
    h('Componentes y balance mensual')
    table([['Actividad','Mensual (%)','Anual (%)']]+[[r['label'].replace('Generación, transmisión','Energía, agua y gas'),f'{r["mom"]:.1f}',f'{r["yoy"]:.1f}'] for r in rows])
    note('Fuente: INEGI, cuadro 1, página 2. Ambas comparaciones utilizan cifras desestacionalizadas.')
    doc.add_picture(str(out.parent/'mom.png'),width=Inches(6.05))
    note('Fuente: INEGI. Comparación respecto al mes inmediato anterior. Las barras no representan aportaciones ponderadas.')
    h('Balance anual por actividad')
    doc.add_picture(str(out.parent/'yoy.png'),width=Inches(6.05))
    note('Fuente: INEGI, cuadro 1. Comparación con el mismo mes del año anterior; cifras desestacionalizadas.')
    h('Cifras originales y acumulado del año')
    p('La lectura de las cifras originales complementa la comparación ajustada. El acumulado compara enero al mes de referencia con el mismo intervalo del año anterior; no es un promedio simple de tasas mensuales.')
    table([['Actividad','Anual original (%)','Enero–mes (%)']]+[[r['label'].replace('Generación, transmisión','Energía, agua y gas'),f'{r["original"]:.1f}',f'{r["ytd"]:.1f}'] for r in rows])
    note('Fuente: INEGI, cuadro 2, página 4. Cifras preliminares. Se conservan las tasas publicadas; no se recalculan con otra versión de los índices.')
    h('Lectura de conjunto')
    positives=sum(r['yoy']>0 for r in parts)
    bullet(f'{positives} de {len(parts)} componentes principales registraron variaciones anuales positivas en cifras desestacionalizadas. Es un conteo no ponderado y no mide su aportación al agregado.')
    bullet(f'La tasa anual original del total fue {total["original"]:.1f}%; el acumulado enero–mes alcanzó {total["ytd"]:.1f}%. Se presentan por separado de las tasas ajustadas.')
    h('Notas metodológicas y de interpretación')
    p('Las series ajustadas eliminan efectos estacionales y de calendario mediante métodos econométricos. Los agregados y sus componentes se ajustan de manera independiente; sus tasas no deben sumarse ni promediarse para reconstruir el total.')
    p('El IMAI abarca el conjunto industrial, no exclusivamente las manufacturas ni la EMIM.' if key=='IMAI' else 'El IGAE aproxima la evolución mensual del sector real y no equivale al PIB trimestral ni incluye la totalidad de las actividades de este último.')
    note(f'Fuente: INEGI, boletín {f["bulletin"]}, publicado el {f["publication_date"]}. Resultados: pp. 1–2; cifras originales: p. 4; ficha metodológica: p. 5. Las cifras están sujetas a revisión. No se atribuyen causas ni récords sin evidencia adicional.')
    note(url);doc.core_properties.title=title+' · '+f['period'];doc.save(out)
