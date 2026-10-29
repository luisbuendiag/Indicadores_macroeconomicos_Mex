import ast,json,shutil
from pathlib import Path
from types import SimpleNamespace
import pytest
from docx import Document
import notas_demanda as N
import notas_demanda_word as V
import build_demand_notes as R
ROOT=Path(__file__).resolve().parents[1]
class PDF:
    def __init__(self,pages):self.pages=[SimpleNamespace(extract_text=lambda p=p:p) for p in pages]
    def __enter__(self):return self
    def __exit__(self,*a):pass

def fixture(monkeypatch,key):
    pages=json.loads((ROOT/f'tests/fixtures/{key.lower()}2026_10_pages.json').read_text())
    monkeypatch.setattr(N.pdfplumber,'open',lambda _:PDF(pages));return pages

def test_consumo_preserves_original_adjusted_and_durability(monkeypatch):
    fixture(monkeypatch,'CONSUMO');f=N.extract(b'fixture','CONSUMO')
    assert f['period']=='Jul 26' and f['index']==113.2 and len(f['originals'])==12
    assert V.row(f,'importado')['yoy']==1.6 and V.row(f,'importado',True)['original']==.1
    assert V.row(f,'total')['mom']==0 and V.row(f,'total',True)['ytd']==2
    assert V.row(f,'nacional_duradero',True)['ytd']==-4 and V.row(f,'importado_duradero',True)['ytd']==12.9
    assert V.row(f,'nacional_servicios')['yoy']==.6 and V.row(f,'nacional_servicios',True)['original']==.7

def test_investment_buyer_not_an_assumption(monkeypatch):
    fixture(monkeypatch,'IMFBCF');f=N.extract(b'fixture','IMFBCF')
    assert f['index']==110.5 and len(f['rows'])==11 and len(f['originals'])==20
    assert V.row(f,'total')['yoy']==6 and V.row(f,'total',True)['original']==6.6
    assert V.row(f,'total_privada',True)['ytd']==.4 and V.row(f,'total_publica',True)['ytd']==10.2
    assert V.row(f,'nacional_transporte')['yoy']==-.8 and V.row(f,'importado_transporte')['yoy']==-7.9
    assert V.row(f,'maquinaria_privada',True)['ytd']==-1.8

@pytest.mark.parametrize('key',['CONSUMO','IMFBCF'])
def test_header_product_period_and_missing_pages_rejected(monkeypatch,key):
    pages=fixture(monkeypatch,key);pages[0]=pages[0].replace('('+N.NAMES[key]+')','(OTRO)')
    with pytest.raises(ValueError):N.extract(b'fixture',key)
    pages=fixture(monkeypatch,key);pages[2]=pages[2].replace('a julio de 2026','a junio de 2026')
    with pytest.raises(ValueError):N.extract(b'fixture',key)
    pages=fixture(monkeypatch,key);pages.pop()
    with pytest.raises(ValueError):N.extract(b'fixture',key)

@pytest.mark.parametrize('key',['CONSUMO','IMFBCF'])
def test_row_order_methodology_and_cover_checked(monkeypatch,key):
    pages=fixture(monkeypatch,key);pages[1]=pages[1].replace('Nacional','Importado')
    with pytest.raises(ValueError):N.extract(b'fixture',key)
    pages=fixture(monkeypatch,key);pages[2]=pages[2].replace('Cifras preliminares','Cifras definitivas')
    with pytest.raises(ValueError):N.extract(b'fixture',key)
    pages=fixture(monkeypatch,key);pages[0]=pages[0].replace('▲ 1.4 %','▲ 5.4 %') if key=='IMFBCF' else pages[0].replace('113.2 0.0 %','113.2 5.0 %')
    with pytest.raises(ValueError):N.extract(b'fixture',key)

def test_investment_duplicates_must_agree(monkeypatch):
    pages=fixture(monkeypatch,'IMFBCF');pages[2]=pages[2].replace('Tipo de bien y de comprador\nFormación bruta de capital fijo 6.6 1.7','Tipo de bien y de comprador\nFormación bruta de capital fijo 6.6 9.7')
    with pytest.raises(ValueError):N.extract(b'fixture','IMFBCF')

def test_consumo_na_not_zero(monkeypatch):
    pages=fixture(monkeypatch,'CONSUMO');pages[2]=pages[2].replace('NA NA','0.0 0.0')
    with pytest.raises(ValueError):N.extract(b'fixture','CONSUMO')

@pytest.mark.parametrize('key',['CONSUMO','IMFBCF'])
def test_word_tables_all_cotejadas_including_hero(monkeypatch,key):
    fixture(monkeypatch,key);f=N.extract(b'fixture',key);d=Document()
    for data in V.tables(f):
        t=d.add_table(rows=0,cols=len(data[0]))
        for row in data:
            for c,v in zip(t.add_row().cells,row):c.text=v
    candidate=SimpleNamespace(tables=d.tables,inline_shapes=[1,2]);V.validate(f,candidate)
    candidate.tables[0].rows[1].cells[0].text='Índice 9999'
    with pytest.raises(ValueError):V.validate(f,candidate)

@pytest.mark.parametrize('key',['CONSUMO','IMFBCF'])
def test_strict_source_allowlist(key):
    directory,name=N.PRODUCTS[key];url=f'https://www.inegi.org.mx/contenidos/saladeprensa/boletines/2026/{directory}/{name}2026_10.pdf'
    assert R.allowed(url,key)
    for bad in [url+'?x=evil',url+'#x',url.replace('www.inegi.org.mx','example.com'),url.replace('https:','http:'),url.replace(name+'2026','otro2026')]:assert not R.allowed(bad,key)
    assert not R.allowed(url,'EMIM')

@pytest.mark.parametrize('case',['disabled','source_error','local_pdf','pause'])
def test_publication_preserves_previous_word(monkeypatch,tmp_path,case):
    for name in ['data/indicadores.json','data/notas/estado.json','config/notas_automaticas.json']:
        p=tmp_path/name;p.parent.mkdir(parents=True,exist_ok=True);shutil.copyfile(ROOT/name,p)
    dest=tmp_path/'downloads/indicadores/CONSUMO/nota/CONSUMO_nota.docx';dest.parent.mkdir(parents=True);dest.write_bytes(b'previous-word')
    cfg=json.loads((tmp_path/'config/notas_automaticas.json').read_text())
    if case=='disabled':
        cfg['enabled']=False;(tmp_path/'config/notas_automaticas.json').write_text(json.dumps(cfg))
        monkeypatch.setattr(R,'download',lambda *a:pytest.fail('No descargar cuando está pausado'))
        assert R.run('CONSUMO',tmp_path,True)['status']=='deshabilitada'
    elif case=='local_pdf':assert R.run('CONSUMO',tmp_path,True,Path('/tmp/local.pdf'))['status']=='error'
    elif case=='source_error':
        monkeypatch.setattr(R,'download',lambda *a:(_ for _ in ()).throw(ValueError('Fuente incorrecta')))
        assert R.run('CONSUMO',tmp_path,True)['status']=='error'
    else:
        cfg['auto_publish']=False;(tmp_path/'config/notas_automaticas.json').write_text(json.dumps(cfg))
        fixture(monkeypatch,'CONSUMO');monkeypatch.setattr(R,'download',lambda *a:b'fixture')
        template=tmp_path/cfg['template'];template.parent.mkdir(parents=True);shutil.copyfile(ROOT/cfg['template'],template)
        for name in ['build_demand_notes.py','notas_demanda.py','notas_demanda_word.py','notas_layout.py','notas_word.py','build_auto_notes.py']:
            p=tmp_path/'scripts'/name;p.parent.mkdir(exist_ok=True);shutil.copyfile(ROOT/'scripts'/name,p)
        monkeypatch.setattr(R.V,'build',lambda *a:shutil.copyfile(template,a[-1]))
        monkeypatch.setattr(R.W,'render',lambda *a:(None,[1,2]));monkeypatch.setattr(R.V,'validate',lambda *a:None)
        monkeypatch.setattr(R.A,'publish',lambda *a:pytest.fail('Publicación deshabilitada'))
        assert R.run('CONSUMO',tmp_path,True)['status']=='validada_sin_publicar'
    assert dest.read_bytes()==b'previous-word'

def test_workflow_python_311_compatibility():
    for name in ['build_demand_notes.py','notas_demanda.py','notas_demanda_word.py','notas_layout.py']:
        ast.parse((ROOT/'scripts'/name).read_text(),feature_version=(3,11))
