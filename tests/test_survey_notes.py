import json
from pathlib import Path
from types import SimpleNamespace
import pytest
import notas_encuestas as N
import notas_encuestas_word as V
import build_survey_notes as R
ROOT=Path(__file__).resolve().parents[1]
class PDF:
    def __init__(self,pages):self.pages=[SimpleNamespace(extract_text=lambda p=p:p) for p in pages]
    def __enter__(self):return self
    def __exit__(self,*a):pass

def fixture(monkeypatch,key):
    pages=json.loads((ROOT/f'tests/fixtures/{key.lower()}2026_10_pages.json').read_text())
    monkeypatch.setattr(N.pdfplumber,'open',lambda _:PDF(pages));return pages

def test_emim_preserves_adjustment_and_complete_sector_universe(monkeypatch):
    pages=fixture(monkeypatch,'EMIM');f=N.extract(b'fixture','EMIM')
    assert f['period']=='Ago 26' and len(f['sectors'])==21
    assert f['rows'][0]['yoy']==4.1 and f['rows'][0]['original']==4.0
    assert f['rows'][1]['yoy']==-1.0 and f['rows'][3]['yoy']==2.7
    assert sum(r['production']>0 for r in f['sectors'])==10
    assert sum(r['production']>0 and r['employment']>0 for r in f['sectors'])==5
    pages[3]=pages[3].replace('339 Otras','338 Otras')
    with pytest.raises(ValueError):N.extract(b'fixture','EMIM')

def test_emoe_distinguishes_units_and_component_series(monkeypatch):
    fixture(monkeypatch,'EMOE');f=N.extract(b'fixture','EMOE')
    assert f['globalrow']['level']==48.4 and f['globalrow']['yoy']==-1.0
    assert f['rows'][0]['level']==48.0 and f['rows'][0]['series']=='desestacionalizada'
    assert f['rows'][1]['series']=='original'
    assert len(f['components'])==20
    c=[r for r in f['components'] if r['sector']=='Construcción']
    assert c[0]['series']=='desestacionalizada' and c[1]['series']=='original'
    c=[r for r in f['components'] if r['sector']=='Manufacturas']
    assert [r['series'] for r in c]==['desestacionalizada']*2+['original']*3

@pytest.mark.parametrize('key',['EMIM','EMOE'])
def test_changed_methodology_product_and_partial_pdf_fail_closed(monkeypatch,key):
    pages=fixture(monkeypatch,key)
    pages[0]=pages[0].replace('('+key+')','(OTRO)')
    with pytest.raises(ValueError):N.extract(b'fixture',key)
    pages=fixture(monkeypatch,key);pages.pop()
    with pytest.raises(ValueError):N.extract(b'fixture',key)

@pytest.mark.parametrize('key,product',[('EMIM','emim'),('EMOE','ice')])
def test_source_allowlist(key,product):
    u=f'https://www.inegi.org.mx/contenidos/saladeprensa/boletines/2026/{product}/{product}2026_10.pdf'
    assert R.allowed(u,key)
    assert not R.allowed(u+'?redirect=evil',key)
    assert not R.allowed(u.replace('www.inegi.org.mx','evil.example'),key)
    assert not R.allowed(u,'IMAI')

def test_word_numeric_tampering_rejected(monkeypatch):
    from docx import Document
    fixture(monkeypatch,'EMOE');f=N.extract(b'fixture','EMOE');d=Document()
    d.add_table(rows=1,cols=2)
    for rows in V.tables(f):
        t=d.add_table(rows=0,cols=len(rows[0]))
        for row in rows:
            for c,value in zip(t.add_row().cells,row):c.text=value
    candidate=SimpleNamespace(tables=d.tables,inline_shapes=[1,2])
    V.validate(f,candidate)
    candidate.tables[1].rows[1].cells[1].text='4800.0'
    with pytest.raises(ValueError):V.validate(f,candidate)

@pytest.mark.parametrize('flag',['enabled','auto_publish'])
def test_configuration_pause_preserves_served_word(monkeypatch,tmp_path,flag):
    import shutil
    for name in ['data/indicadores.json','data/notas/estado.json','config/notas_automaticas.json']:
        p=tmp_path/name;p.parent.mkdir(parents=True,exist_ok=True);shutil.copyfile(ROOT/name,p)
    cfg=json.loads((tmp_path/'config/notas_automaticas.json').read_text());cfg[flag]=False
    (tmp_path/'config/notas_automaticas.json').write_text(json.dumps(cfg))
    dest=tmp_path/'downloads/indicadores/EMIM/nota/EMIM_nota.docx';dest.parent.mkdir(parents=True);dest.write_bytes(b'previous-word')
    if flag=='enabled':
        monkeypatch.setattr(R,'download',lambda *a:pytest.fail('No descargar si generación pausada'))
        assert R.run('EMIM',tmp_path,True)['status']=='deshabilitada'
    else:
        fixture(monkeypatch,'EMIM');monkeypatch.setattr(R,'download',lambda *a:b'fixture')
        template=cfg['template'];p=tmp_path/template;p.parent.mkdir(parents=True,exist_ok=True);shutil.copyfile(ROOT/template,p)
        for name in ['build_survey_notes.py','notas_encuestas.py','notas_encuestas_word.py','notas_word.py','build_auto_notes.py']:
            p=tmp_path/'scripts'/name;p.parent.mkdir(exist_ok=True);shutil.copyfile(ROOT/'scripts'/name,p)
        monkeypatch.setattr(R.V,'build',lambda *a:shutil.copyfile(ROOT/cfg['template'],a[-1]))
        monkeypatch.setattr(R.W,'render',lambda *a:(None,[1,2]))
        monkeypatch.setattr(R.V,'validate',lambda *a:None)
        monkeypatch.setattr(R.A,'publish',lambda *a:pytest.fail('No publicar con configuración pausada'))
        assert R.run('EMIM',tmp_path,True)['status']=='validada_sin_publicar'
    assert dest.read_bytes()==b'previous-word'
