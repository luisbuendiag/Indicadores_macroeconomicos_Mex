import ast,json,shutil
from pathlib import Path
from types import SimpleNamespace
import pytest
from docx import Document
import notas_precios as N
import notas_precios_word as V
import build_price_notes as R
from sources import inegi_inpp as P
ROOT=Path(__file__).resolve().parents[1]
class Words:
    def __init__(self,words):self.words=words
    def extract_words(self):return self.words
class Page:
    def __init__(self,text,generic):self.text=text;self.g=generic;self.width=generic['width'];self.height=generic['height']
    def extract_text(self):return self.text
    def extract_words(self):return [{'text':'1/','top':self.g['foot']}]
    def crop(self,box):return Words(self.g['columns'][0 if box[0]==0 else 1])
class PDF:
    def __init__(self,fixture):self.pages=[Page(t,fixture['generic']) for t in fixture['pages']]
    def __enter__(self):return self
    def __exit__(self,*a):pass

def fixture(monkeypatch,key):
    f=json.loads((ROOT/f'tests/fixtures/{key.lower()}2026_10_precios.json').read_text())
    monkeypatch.setattr(N.pdfplumber,'open',lambda _:PDF(f));return f

def test_inpc_rates_are_not_incidences_or_quincenal(monkeypatch):
    fixture(monkeypatch,'INPC');f=N.extract(b'fixture','INPC');r=f['rows'][0]
    assert f['period']=='Sep 26' and f['index']==146.075
    assert r['mom']==.42 and r['yoy']==3.45 and r['prior_yoy']==3.76
    assert r['incidence']==.421 and r['annual_incidence']==3.455
    assert len(f['rows'])==16 and len(f['generics'])==20
    assert f['generics'][0]['label']=='Jitomate' and f['generics'][0]['incidence']==.133
    assert f['generics'][12]['label']=='Paquetes de internet, telefonía y televisión de paga'

def test_inpp_coverage_and_destination_separated(monkeypatch):
    fixture(monkeypatch,'INPP');f=N.extract(b'fixture','INPP')
    assert f['rows'][0]['mom']==1.37 and f['rows'][0]['yoy']==4.05
    assert [g[0]['yoy'] for g in f['groups']]==[4.05,4.25,3.97]
    assert f['destinations'][0]['incidence']==1.137 and f['destinations'][0]['mom']==1.14
    assert len(f['generics'])==30 and f['generics'][0]['incidence']==.587
    assert f['index'] is None
    assert V.selected(f)[3]['label']=='Aguacate'

@pytest.mark.parametrize('key',['INPC','INPP'])
def test_wrong_product_month_and_columns_fail_closed(monkeypatch,key):
    f=fixture(monkeypatch,key);f['pages'][0]=f['pages'][0].replace('('+key+')','(OTRO)')
    with pytest.raises(ValueError):N.extract(b'fixture',key)
    f=fixture(monkeypatch,key);f['pages'][1]=f['pages'][1].replace('2024 2025 2026','2025 2024 2026')
    with pytest.raises(ValueError):N.extract(b'fixture',key)
    f=fixture(monkeypatch,key);f['pages'][1]=f['pages'][1].replace('Actividades primarias','Otra actividad') if key=='INPP' else f['pages'][1].replace('Subyacente','Otra actividad')
    with pytest.raises(ValueError):N.extract(b'fixture',key)

@pytest.mark.parametrize('key',['INPC','INPP'])
def test_wrong_sign_or_missing_generic_is_rejected(monkeypatch,key):
    f=fixture(monkeypatch,key)
    for w in f['generic']['columns'][0]:
        if w['text'] in ['0.133','0.587']:w['text']='-0.133'
    with pytest.raises(ValueError):N.extract(b'fixture',key)
    f=fixture(monkeypatch,key);f['generic']['columns'][1]=[w for w in f['generic']['columns'][1] if w['text']!='-0.026' and w['text']!='-0.056']
    with pytest.raises(ValueError):N.extract(b'fixture',key)

@pytest.mark.parametrize('key',['INPC','INPP'])
def test_every_word_table_including_hero_is_cotejada(monkeypatch,key):
    fixture(monkeypatch,key);f=N.extract(b'fixture',key);d=Document()
    for rows in V.tables(f):
        t=d.add_table(rows=0,cols=len(rows[0]))
        for row in rows:
            for cell,v in zip(t.add_row().cells,row):cell.text=v
    candidate=SimpleNamespace(tables=d.tables,inline_shapes=[1,2]);V.validate(f,candidate)
    candidate.tables[2].rows[1].cells[2].text='133.000'
    with pytest.raises(ValueError):V.validate(f,candidate)

@pytest.mark.parametrize('period,urlpart',[('Sep 26','2026/inpp/inpp2026_10.pdf'),('Dic 26','2027/inpp/inpp2027_01.pdf')])
def test_inpp_filename_is_publication_month(period,urlpart):
    assert P.expected_bulletin_url(period).endswith(urlpart)

def test_discovery_verifies_pdf_reference_not_just_url(monkeypatch):
    text=json.loads((ROOT/'tests/fixtures/inpp2026_10_precios.json').read_text())['pages'][0]
    calls=[]
    monkeypatch.setattr(P,'_first_page',lambda url:(calls.append(url),text)[1])
    assert P._discover_bulletin('Sep 26')['url'].endswith('2026_10.pdf')
    assert P._discover_bulletin('Ago 26') is None
    assert calls[1].endswith('2026_09.pdf')

@pytest.mark.parametrize('key,url',[('INPC','https://www.inegi.org.mx/contenidos/saladeprensa/boletines/2026/inpc/inpc_2q2026_10.pdf'),('INPP','https://www.inegi.org.mx/contenidos/saladeprensa/boletines/2026/inpp/inpp2026_10.pdf')])
def test_urls_are_exact_products_and_monthly_only(key,url):
    assert R.allowed(url,key)
    assert not R.allowed(url+'?x=evil',key)
    assert not R.allowed(url.replace('www.inegi.org.mx','evil.example'),key)
    assert not R.allowed(url.replace('inpc_2q','inpc_1q'),key) if key=='INPC' else not R.allowed(url,'INPC')

def test_source_change_rolls_back_if_publication_fails(monkeypatch,tmp_path):
    path=tmp_path/'data/indicadores.json';path.parent.mkdir(parents=True)
    before=json.dumps({'indicators':{'INPP':{'url_boletin_oficial':'old','last_observation':'Sep 26'}}}).encode();path.write_bytes(before)
    monkeypatch.setattr(R.A,'publish',lambda *a:(_ for _ in ()).throw(ValueError('Fallo controlado')))
    facts={'indicator':'INPP','period':'Sep 26','bulletin':'620/26'}
    with pytest.raises(ValueError):R.publish(tmp_path,None,facts,None,'new',None,None,'old')
    assert path.read_bytes()==before
    path.write_text(json.dumps({'indicators':{'INPP':{'url_boletin_oficial':'changed','last_observation':'Sep 26'}}}))
    with pytest.raises(ValueError):R.publish(tmp_path,None,facts,None,'new',None,None,'old')
    assert 'changed' in path.read_text()

@pytest.mark.parametrize('case',['disabled','local_pdf','download_error'])
def test_failures_and_pause_preserve_current_word(monkeypatch,tmp_path,case):
    for name in ['data/indicadores.json','data/notas/estado.json','config/notas_automaticas.json']:
        p=tmp_path/name;p.parent.mkdir(parents=True,exist_ok=True);shutil.copyfile(ROOT/name,p)
    dest=tmp_path/'downloads/indicadores/INPP/nota/INPP_nota.docx';dest.parent.mkdir(parents=True);dest.write_bytes(b'previous-word')
    if case=='disabled':
        p=tmp_path/'config/notas_automaticas.json';cfg=json.loads(p.read_text());cfg['enabled']=False;p.write_text(json.dumps(cfg))
        monkeypatch.setattr(R,'download',lambda *a:pytest.fail('Generación pausada'))
        assert R.run('INPP',tmp_path,True)['status']=='deshabilitada'
    elif case=='local_pdf':assert R.run('INPP',tmp_path,True,Path('/tmp/a.pdf'))['status']=='error'
    else:
        monkeypatch.setattr(R,'download',lambda *a:(_ for _ in ()).throw(ValueError('PDF no válido')))
        assert R.run('INPP',tmp_path,True)['status']=='error'
    assert dest.read_bytes()==b'previous-word'

def test_python_311_compatibility():
    for name in ['notas_precios.py','notas_precios_word.py','build_price_notes.py','sources/inegi_inpp.py']:
        ast.parse((ROOT/'scripts'/name).read_text(),feature_version=(3,11))

def test_source_update_uses_validated_period_without_changing_observations(monkeypatch,tmp_path):
    path=tmp_path/'data/indicadores.json';path.parent.mkdir(parents=True)
    payload={'indicators':{'INPP':{'url_boletin_oficial':'old','last_observation':'Sep 26','observations':[{'values':[4.05]}]}}};path.write_text(json.dumps(payload))
    monkeypatch.setattr(R.A,'publish',lambda *a:{'status':'vigente'})
    f={'indicator':'INPP','period':'Sep 26','bulletin':'620/26'}
    assert R.publish(tmp_path,None,f,None,'new',None,None,'old')['status']=='vigente'
    p=json.loads(path.read_text())['indicators']['INPP']
    assert p['url_boletin_oficial']=='new' and p['boletin_validado'] is True
    assert p['observations']==payload['indicators']['INPP']['observations']
