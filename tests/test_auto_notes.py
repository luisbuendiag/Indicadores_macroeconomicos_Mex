"""Publicación transaccional, revisiones, fallos y revisión obligatoria del piloto."""
import copy
import hashlib
import io
import json
import shutil
import urllib.error
from pathlib import Path
from types import SimpleNamespace

import pytest
from docx import Document

import build_auto_notes as A
import lib_notes_state as S
import notas_imoam as I
import notas_review as R

ROOT=Path(__file__).resolve().parents[1]
URL='https://www.inegi.org.mx/contenidos/saladeprensa/boletines/2026/imoam/imoam2026_10.pdf'

@pytest.fixture
def facts():
    return dict(indicator='IMOAM',ym='2026-09',period='Sep 26',publication_date='2026-10-28',bulletin='639/26',
                values=[108,105.1,111.1,.019,-.009,.047],base=None,pages=['Cuadro 1','Metodología completa'],
                source_sha256=hashlib.sha256(b'boletin').hexdigest(),evidence=['Índice 108.0 105.1 111.1','Variación anual 1.9% -0.9% 4.7%'],series='original',statistical_status='estimado')

@pytest.fixture
def sandbox(tmp_path,monkeypatch,facts):
    root=tmp_path
    for folder in ('config','scripts','data','downloads/indicadores/IMOAM/nota'):(root/folder).mkdir(parents=True,exist_ok=True)
    cfg=dict(template='template.docx',rules='config/reglas.md',max_source_pages=12,max_note_pages=2)
    (root/'config/notas_automaticas.json').write_text(json.dumps(cfg))
    (root/'config/reglas.md').write_text('Reglas de precisión y Word institucional')
    for name in ('build_auto_notes.py','notas_imoam.py','notas_word.py','notas_review.py'):
        shutil.copyfile(ROOT/'scripts'/name,root/'scripts'/name)
    Document().save(root/'template.docx')
    old=Document();old.add_paragraph('Nota anterior');old.save(root/'downloads/indicadores/IMOAM/nota/IMOAM_nota.docx')
    ind=dict(key='IMOAM',last_observation='Sep 26',url_boletin_oficial=URL,
             observations=[dict(period='Sep 26',values=facts['values'].copy())],
             nota={'available':True,'format':'docx','path':'downloads/indicadores/IMOAM/nota/IMOAM_nota.docx','periodo':'Ago-26','automatic':False},
             nota_disponible=True)
    (root/'data/indicadores.json').write_text(json.dumps({'indicators':{'IMOAM':ind,'OTHER':{'untouched':True}}}))
    monkeypatch.delenv('OPENAI_API_KEY',raising=False);monkeypatch.setenv('NOTES_REVIEW_MODEL',R.VERSION)
    monkeypatch.setenv('NOTES_AUTO_PUBLISH','true');monkeypatch.delenv('NOTES_WORK_DIR',raising=False)
    monkeypatch.setattr(A,'download',lambda url:b'boletin')
    monkeypatch.setattr(I,'extract',lambda pdf:copy.deepcopy(facts))
    def build(facts,url,template,out):
        doc=Document();doc.add_paragraph('Candidata nueva '+facts['period']);doc.save(out);return 'Texto de la nota'
    monkeypatch.setattr(A.W,'build',build)
    def render(out,max_pages):
        pdf=out.with_suffix('.pdf');pdf.write_bytes(b'PDF');png=out.parent/'nota-pagina-1.png';png.write_bytes(b'PNG');return pdf,[png]
    monkeypatch.setattr(A.W,'render',render)
    def raster(pdf,prefix):
        png=prefix.with_suffix('.png');png.write_bytes(b'PNG');return [png]
    monkeypatch.setattr(A.W,'raster',raster)
    return root


def approve(*args):
    return {'passed':True,'result':dict(datos=True,logica=True,estructura=True,render=True,issues=[]),'model':R.VERSION,'kind':'reglas_python'}


def old_doc(root):return (root/'downloads/indicadores/IMOAM/nota/IMOAM_nota.docx').read_bytes()


def test_no_credentials_local_checks_can_publish(sandbox,monkeypatch):
    monkeypatch.delenv('NOTES_REVIEW_MODEL',raising=False)
    monkeypatch.setenv('OPENAI_API_KEY','unused-placeholder')
    monkeypatch.setattr(A.R,'review',approve)
    def no_external_call(*a,**k):pytest.fail('No debe llamar a una API')
    monkeypatch.setattr(urllib.request,'urlopen',no_external_call)
    old=old_doc(sandbox)
    assert A.run(sandbox,True,True)['status']=='vigente' and old_doc(sandbox)!=old


def test_success_is_atomic_traceable_and_not_repeated(sandbox):
    old=old_doc(sandbox);state=A.run(sandbox,True,True,reviewer=approve)
    assert state['status']=='vigente' and old_doc(sandbox)!=old
    pub=state['published'];archive=sandbox/pub['archive']
    assert all((archive/x).is_file() for x in ('boletin.pdf','hechos.json','revision.json','trazabilidad.csv','IMOAM_nota.docx'))
    assert S.digest(sandbox/'downloads/indicadores/IMOAM/nota/IMOAM_nota.docx')==pub['docx_sha256']
    def unexpected(*a):pytest.fail('Revisión repetida para el mismo boletín')
    assert A.run(sandbox,True,True,reviewer=unexpected)['unchanged']


def test_same_period_revision_and_rule_change_regenerate(sandbox,monkeypatch):
    first=A.run(sandbox,True,True,reviewer=approve)['published']['fingerprint']
    monkeypatch.setattr(A,'download',lambda url:b'boletin revisado')
    second=A.run(sandbox,True,True,reviewer=approve)['published']['fingerprint']
    assert second!=first
    (sandbox/'config/reglas.md').write_text('Reglas revisadas')
    assert A.run(sandbox,True,True,reviewer=approve)['published']['fingerprint']!=second


def test_changed_period_must_match_dashboard(sandbox,facts,monkeypatch):
    old=old_doc(sandbox);facts.update(ym='2026-10',period='Oct 26')
    state=A.run(sandbox,True,True,reviewer=approve)
    assert state['status']=='error' and old_doc(sandbox)==old


def test_revision_discrepancy_does_not_publish(sandbox,facts):
    old=old_doc(sandbox);facts['values'][0]=109
    assert A.run(sandbox,True,True,reviewer=approve)['status']=='error'
    assert old_doc(sandbox)==old


def test_render_failure_keeps_previous(sandbox,monkeypatch):
    old=old_doc(sandbox)
    def fail(*a):raise ValueError('Página desbordada')
    monkeypatch.setattr(A.W,'render',fail)
    assert A.run(sandbox,True,True,reviewer=approve)['status']=='error'
    assert old_doc(sandbox)==old


def test_qa_rejection_preserves_word_and_review_evidence(sandbox):
    old=old_doc(sandbox)
    def reject(*a):return {'passed':False,'result':dict(datos=True,logica=True,estructura=True,render=False,issues=['Texto encimado']),'model':R.VERSION,'kind':'reglas_python'}
    assert A.run(sandbox,True,True,reviewer=reject)['status']=='rechazada'
    assert old_doc(sandbox)==old and list((sandbox/'.notes-work').rglob('revision.json'))


def test_publication_needs_explicit_switch(sandbox,monkeypatch):
    old=old_doc(sandbox);monkeypatch.setenv('NOTES_AUTO_PUBLISH','false')
    assert A.run(sandbox,True,True,reviewer=approve)['status']=='validada_sin_publicar'
    assert old_doc(sandbox)==old


def test_published_metadata_survives_next_provision_and_detects_tamper(sandbox):
    A.run(sandbox,True,True,reviewer=approve)
    ind=json.loads((sandbox/'data/indicadores.json').read_text())['indicators']['IMOAM']
    ind['nota']['automatic']=False;ind['nota']['periodo']='Ago-26'
    S.decorate(ind,sandbox)
    assert ind['nota']['automatic'] and ind['nota']['periodo']=='Sep 26'
    path=sandbox/'downloads/indicadores/IMOAM/nota/IMOAM_nota.docx';path.write_bytes(b'corrupto')
    S.decorate(ind,sandbox)
    assert ind['nota']['status']=='error' and not ind['nota']['automatic']


def test_previous_auto_note_is_marked_pending_on_new_period(sandbox):
    A.run(sandbox,True,True,reviewer=approve)
    ind=json.loads((sandbox/'data/indicadores.json').read_text())['indicators']['IMOAM'];ind['last_observation']='Oct 26'
    S.decorate(ind,sandbox)
    assert ind['nota']['status']=='pendiente' and ind['nota']['periodo']=='Sep 26'


def test_failed_publication_rolls_back_service_files(sandbox,monkeypatch):
    old=old_doc(sandbox);before=(sandbox/'data/indicadores.json').read_bytes();atomic=A.atomic;failed=False
    def fail_once(path,data):
        nonlocal failed
        if path==sandbox/'data/indicadores.json' and not failed:
            failed=True;raise OSError('disco')
        return atomic(path,data)
    monkeypatch.setattr(A,'atomic',fail_once)
    assert A.run(sandbox,True,False,reviewer=approve)['status']=='error'
    assert old_doc(sandbox)==old and (sandbox/'data/indicadores.json').read_bytes()==before

@pytest.mark.parametrize('url',[URL.replace('www.inegi.org.mx','evil.example'),URL.replace('/imoam/','/ioae/'),URL.replace('https:','http:'),'https://www.inegi.org.mx@evil.example/x.pdf'])
def test_reject_wrong_source(url):assert not A.official_url(url)


def test_uncertainty_and_negative_results_are_not_hidden(facts):
    text=' '.join(p for _,p in I.paragraphs(facts))
    assert 'cero y valores negativos' in text and 'sin ajuste estacional' in text
    facts['values'][3:]=[-.019,-.047,-.009]
    text=' '.join(p for _,p in I.paragraphs(facts))
    assert 'una disminución' in text and 'terreno negativo' in text and 'un aumento' not in text

@pytest.mark.parametrize('result',[{},dict(datos='true',logica=True,estructura=True,render=True,issues=[]),dict(datos=True,logica=True,estructura=True,render=True,issues='ok')])
def test_qa_schema_must_be_complete_and_typed(result):
    with pytest.raises(ValueError):R.validate(result)



def test_rejected_input_waits_for_intervention(sandbox):
    def reject(*a):return {'passed':False,'result':dict(datos=False,logica=True,estructura=True,render=True,issues=['Fuente ambigua']),'model':R.VERSION,'kind':'reglas_python'}
    A.run(sandbox,True,True,reviewer=reject)
    def unexpected(*a):pytest.fail('No debe repetir una revisión rechazada sin cambios')
    assert A.run(sandbox,True,True,reviewer=unexpected)['unchanged']
    assert A.run(sandbox,True,True,reviewer=approve,retry_review=True)['status']=='vigente'


def test_transport_failures_have_bounded_retries(sandbox):
    count=0
    def fail(*a):
        nonlocal count
        count+=1;raise TimeoutError('renderer')
    for _ in range(5):A.run(sandbox,True,True,reviewer=fail)
    assert count==3


class PDF:
    def __init__(self,texts):self.pages=[SimpleNamespace(extract_text=lambda text=t:text) for t in texts]
    def __enter__(self):return self
    def __exit__(self,*a):pass


def test_complete_official_bulletin_is_traced_without_assumed_base(monkeypatch):
    pages=json.loads((ROOT/'tests/fixtures/imoam2026_10_pages.json').read_text())
    monkeypatch.setattr(I.pdfplumber,'open',lambda b:PDF(pages))
    facts=I.extract(b'%PDF fixture')
    assert facts['values']==pytest.approx([108,105.1,111.1,.019,-.009,.047])
    assert len(facts['pages'])==4 and 'FICHA METODOLÓGICA' in facts['pages'][3]
    assert facts['base'] is None
    trace=I.trace(facts,URL)
    assert len(trace)==14 and all(r['ubicacion']=='página 1, cuadro 1' for r in trace[:8])
    assert facts['context']['coverage']==90
    assert facts['context']['evaluation_start']=='2019-01'
    assert facts['context']['evaluation_end']=='2026-08'
    assert facts['context']['days_ahead']==18 and facts['context']['imai_lag']==40
    assert trace[6]['valor']=='6.0' and trace[7]['valor']=='5.6'
    assert trace[3]['calculo']=='valor porcentual / 100'

@pytest.mark.parametrize('change',['method','series','blank'])
def test_methodological_change_or_unreadable_page_requires_review(monkeypatch,change):
    pages=json.loads((ROOT/'tests/fixtures/imoam2026_10_pages.json').read_text())
    if change=='method':pages=[p.replace('estadística experimental','estadística distinta') for p in pages]
    if change=='series':pages[0]=pages[0].replace('cifras originales','cifras ajustadas')
    if change=='blank':pages[3]=''
    monkeypatch.setattr(I.pdfplumber,'open',lambda b:PDF(pages))
    with pytest.raises(ValueError):I.extract(b'%PDF fixture')


def test_automatic_flag_requires_reproducible_audit(sandbox):
    A.run(sandbox,True,True,reviewer=approve)
    ind=json.loads((sandbox/'data/indicadores.json').read_text())['indicators']['IMOAM']
    assert not S.validate_published(ind,sandbox)
    pub=S.load_state(sandbox/'data/notas/estado.json')['IMOAM']['published']
    source=sandbox/pub['archive']/'boletin.pdf';source.write_bytes(b'fuente cambiada')
    assert S.validate_published(ind,sandbox)


def test_forged_automatic_flag_is_rejected(sandbox):
    ind=json.loads((sandbox/'data/indicadores.json').read_text())['indicators']['IMOAM']
    ind['nota']['automatic']=True
    assert S.validate_published(ind,sandbox)


def test_local_pdf_cannot_bypass_official_source_publication(sandbox,tmp_path):
    local=tmp_path/'input.pdf';local.write_bytes(b'boletin')
    old=old_doc(sandbox)
    assert A.run(sandbox,True,True,pdf_path=local,reviewer=approve)['status']=='error'
    assert old_doc(sandbox)==old


def test_year_change_requires_updated_institutional_template(sandbox):
    p=sandbox/'config/notas_automaticas.json';c=json.loads(p.read_text());c['template_year']=2025;p.write_text(json.dumps(c))
    assert A.run(sandbox,True,True,reviewer=approve)['status']=='error'


def test_note_notice_renders_period_and_pending_message():
    node=shutil.which('node')
    if not node:pytest.skip('Node no disponible')
    script=r'''
const fs=require('node:fs'),assert=require('node:assert/strict');
const src=fs.readFileSync('assets/js/app.js','utf8');
const part=src.split('  if (ind.nota?.status) {')[1].split('\n  }')[0];
const render=new Function('ind','panel','el',part);
const children=[];
render({nota:{status:'pendiente',periodo:'Ago-26',message:'Actualizando nota'}},{append:x=>children.push(x)},(tag,attrs,text)=>({tag,attrs,text}));
assert.equal(children.length,1);assert.equal(children[0].attrs.role,'status');
assert.ok(children[0].text.includes('Ago-26'));assert.ok(children[0].text.includes('Actualizando nota'));
'''
    import subprocess
    subprocess.run([node,'-e',script],cwd=ROOT,check=True,capture_output=True,text=True)


def test_missing_evaluation_requires_review():
    pages=json.loads((ROOT/'tests/fixtures/imoam2026_10_pages.json').read_text())
    pages[2]=pages[2].replace('90 %','noventa por ciento')
    with pytest.raises(ValueError): I.context_from_pages(pages[:3])
    with pytest.raises(ValueError): I.context_from_pages(['','','',''])


def test_two_pages_required_before_publication(sandbox):
    cfg=json.loads((sandbox/'config/notas_automaticas.json').read_text())
    cfg['min_note_pages']=2
    (sandbox/'config/notas_automaticas.json').write_text(json.dumps(cfg))
    old=old_doc(sandbox)
    result=A.run(sandbox,True,True,reviewer=approve)
    assert result['status']=='error' and old_doc(sandbox)==old
    assert 'dos páginas' in result['message']


def test_publication_defaults_to_config_without_variables(sandbox,monkeypatch):
    monkeypatch.delenv('NOTES_AUTO_PUBLISH',raising=False)
    cfg=json.loads((sandbox/'config/notas_automaticas.json').read_text());cfg['auto_publish']=True
    (sandbox/'config/notas_automaticas.json').write_text(json.dumps(cfg))
    assert A.run(sandbox,True,True,reviewer=approve)['status']=='vigente'


def test_no_paid_service_in_production_path():
    src=(ROOT/'scripts/notas_review.py').read_text()
    flow=(ROOT/'.github/workflows/update-data.yml').read_text()
    assert 'urllib' not in src and 'api.openai.com' not in src
    assert 'OPENAI_API_KEY' not in flow and 'NOTES_REVIEW_MODEL' not in flow


def test_changed_methodology_requires_intervention():
    pages=json.loads((ROOT/'tests/fixtures/imoam2026_10_pages.json').read_text())
    pages[3]=pages[3].replace('CENACE','NUEVA FUENTE')
    with pytest.raises(ValueError,match='metodológica'):I.context_from_pages(pages)

@pytest.fixture
def local_review_case(tmp_path,monkeypatch):
    from PIL import Image
    pages=json.loads((ROOT/'tests/fixtures/imoam2026_10_pages.json').read_text())
    monkeypatch.setattr(I.pdfplumber,'open',lambda b:PDF(pages))
    facts=I.extract(b'fixture')
    I.write_evidence(tmp_path,facts,b'fixture',URL)
    monkeypatch.setattr(I,'extract',lambda pdf:copy.deepcopy(facts))
    shutil.copyfile(ROOT/'docs/examples/notas/IMOAM_nota_piloto.docx',tmp_path/'IMOAM_nota.docx')
    doc=Document(tmp_path/'IMOAM_nota.docx')
    from docx.oxml.ns import qn
    draft=' '.join(t.text or '' for t in doc.element.body.iter(qn('w:t')))
    images=[]
    for label,count in [('Boletín oficial página',4),('Nota Word página',2)]:
        for i in range(count):
            p=tmp_path/(label.split()[0]+str(i)+'.png');Image.new('RGB',(1020,1320),'white').save(p);images.append((label,p))
    (tmp_path/'IMOAM_nota.pdf').write_bytes(b'pdf-placeholder')
    # Geometría simulada; el renderer real se comprueba en la ejecución integrada.
    text=draft+' texto'*300
    pdf=SimpleNamespace(pages=[SimpleNamespace(extract_text=lambda:text,chars=[],width=612,height=792) for _ in range(2)])
    class RenderPDF:
        def __enter__(self):return pdf
        def __exit__(self,*a):pass
    monkeypatch.setattr(R.pdfplumber,'open',lambda path:RenderPDF())
    return tmp_path,facts,draft,images


def test_local_validator_accepts_complete_candidate(local_review_case):
    work,facts,draft,images=local_review_case
    result=R.review(facts,draft,'reglas',images)
    assert result['passed'] and result['kind']=='reglas_python'
    assert result['limits'] and 'visual' not in result['result']


@pytest.mark.parametrize('damage',['table','trace','page','bullets'])
def test_local_validator_blocks_corrupted_candidate(local_review_case,damage):
    work,facts,draft,images=local_review_case
    if damage=='table':
        d=Document(work/'IMOAM_nota.docx');d.tables[1].cell(1,1).text='999.0';d.save(work/'IMOAM_nota.docx')
    if damage=='trace':
        p=work/'trazabilidad.csv';p.write_text(p.read_text(encoding='utf-8-sig').replace('108.0','999.0'),encoding='utf-8-sig')
    if damage=='page':images=images[:-1]
    if damage=='bullets':
        d=Document(work/'IMOAM_nota.docx')
        from docx.oxml.ns import qn
        for n in d.element.body.findall('.//'+qn('w:numPr')):n.getparent().remove(n)
        d.save(work/'IMOAM_nota.docx')
    result=R.review(facts,draft,'reglas',images)
    assert not result['passed'] and result['result']['issues']
